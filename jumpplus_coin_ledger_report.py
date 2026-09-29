"""ジャンプ＋コイン出納レポート (report_id: jumpplus-coin-ledger).

One run produces three XLSX files that are delivered together as a single
OTP-protected multi-file delivery (see docs/jumpplus-coin-ledger-report.md):

- App  : 【少年ジャンプ＋】消費コイン_yyyy年mm月期_yymmdd.xlsx      (iOS + And)
- WEB  : 【少年ジャンプ＋】消費コイン_WEB_yyyy年mm月期_yymmdd.xlsx  (Web)
- 話売商品: 話売商品_一覧_yymm.xlsx                                  (current episode master)

External communication performed by this module (all via the runtime
service account / ADC, nothing else):
- BigQuery (project BIGQUERY_PROJECT_ID, read-only SELECTs on the three
  source tables below)
- Google Drive (template download, generated-file upload to the official
  folder, best-effort trash of this run's own uploads on partial failure)
- Google Cloud Storage (generated-file upload to BUCKET_NAME for OTP
  delivery)

Deliberately never reads `report_plus_monthly_coin_content_report_mom`: that
table is fixed to the previous month, so any target_month other than "last
month" would silently read the wrong data. The content query below reads the
base table filtered by `purchase_date_month_jst = @target_month` instead, and
`_checked_table()` refuses any configured table name ending in `_mom`.
"""

from __future__ import annotations

import json
import logging
import os
import secrets
import sys
import tempfile
import time
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta, timezone
from pathlib import Path
from typing import Any, Callable
from zoneinfo import ZoneInfo

from google.cloud import bigquery

# ---------------------------------------------------------------------------
# Report identity / defaults
# ---------------------------------------------------------------------------

REPORT_ID = "jumpplus-coin-ledger"
REPORT_NAME = "ジャンプ＋コイン出納レポート"
DELIVERY_CUSTOMER_NAME = "集英社 少年ジャンプ＋"

DEFAULT_LEDGER_TABLE = "jumpplus-4a5f4.dataset_datamart_tables.report_plus_monthly_coin_report"
DEFAULT_CONTENT_TABLE = "jumpplus-4a5f4.dataset_datamart_tables.report_plus_monthly_coin_content_report"
DEFAULT_PRODUCT_MASTER_TABLE = "jumpplus-4a5f4.dataset_aggregation_tables.raise_master_contents_works"

# Official Drive folder / templates placed by システム管理ユニット (2026-09-28).
DEFAULT_OUTPUT_FOLDER_ID = "13M5HROiHR9mHGOjQ_a5mmJ8eTNWrVATA"
DEFAULT_GCS_PREFIX = "reports/jumpplus-coin-ledger"

# Same four domains the Admin UI applies when its allowed-domains field is
# left blank (app.py DEFAULT_ALLOWED_DOMAINS) -- the existing ICE Report
# Generator standard, not a report-specific allowlist.
DEFAULT_ALLOWED_DOMAINS = ("shueisha.co.jp", "sur.co.jp", "hitotsubashi.co.jp", "impress.co.jp")

COIN_TYPES = (
    "pay_coin",
    "pay_bonus_coin",
    "free_ad_coin",
    "free_bonus_coin",
    "pay_gift_coin",
    "reward_video_ad_coin",
)
LEDGER_PLATFORMS = ("iOS", "And", "Web")
CONTENT_APP_IDS = {31: "iOS", 32: "And", 101: "Web"}
PURCHASE_TYPES = ("episode", "book")


@dataclass(frozen=True)
class OutputFileSpec:
    file_key: str
    label: str
    template_env: str
    default_template_file_id: str


OUTPUT_FILES: tuple[OutputFileSpec, ...] = (
    OutputFileSpec("app", "App版", "JUMPPLUS_COIN_LEDGER_APP_TEMPLATE_FILE_ID", "1jZuoDggHBtfeiEPrRscImSo_k--l_I66"),
    OutputFileSpec("web", "WEB版", "JUMPPLUS_COIN_LEDGER_WEB_TEMPLATE_FILE_ID", "1kD_KjDSNmhp-JrP8OkAUaDosobeUvftB"),
    OutputFileSpec(
        "product", "話売商品一覧", "JUMPPLUS_COIN_LEDGER_PRODUCT_TEMPLATE_FILE_ID", "1eMWR4uh8y06-He6I0_7csTy7EStbztMt"
    ),
)
OUTPUT_FILE_KEYS = tuple(spec.file_key for spec in OUTPUT_FILES)


class CoinLedgerReportError(Exception):
    def __init__(
        self,
        code: str,
        message: str | None = None,
        *,
        status_code: int = 400,
        retryable: bool = False,
        **details: Any,
    ) -> None:
        super().__init__(message or code)
        self.code = code
        self.status_code = status_code
        self.retryable = retryable
        self.details = details


# ---------------------------------------------------------------------------
# Logging
#
# The app's root logger has no level configured, so plain logging.info()
# never reaches Cloud Logging (see tests/test_ad_revenue_scheduler_
# observability.py). Rather than promoting operational lines to WARNING like
# the ad-revenue report had to, this report logs through its own logger that
# writes one structured JSON line per record to stdout, which Cloud Run's
# logging agent parses into a real Cloud Logging severity (INFO stays INFO).
# Only identifiers, statuses and counts are ever passed as fields -- never
# SQL text, cell values, tokens, PINs, signed URLs or email addresses.
# ---------------------------------------------------------------------------

LOGGER_NAME = "ice_report.jumpplus_coin_ledger"


class _CloudLoggingJsonFormatter(logging.Formatter):
    def format(self, record: logging.LogRecord) -> str:
        payload: dict[str, Any] = {
            "severity": record.levelname,
            "message": record.getMessage(),
            "component": REPORT_ID,
        }
        fields = getattr(record, "fields", None)
        if isinstance(fields, dict):
            payload.update(fields)
        return json.dumps(payload, ensure_ascii=False, default=str)


def _build_logger() -> logging.Logger:
    built = logging.getLogger(LOGGER_NAME)
    if not any(getattr(h, "_coin_ledger_handler", False) for h in built.handlers):
        handler = logging.StreamHandler(sys.stdout)
        handler.setFormatter(_CloudLoggingJsonFormatter())
        handler._coin_ledger_handler = True  # type: ignore[attr-defined]
        built.addHandler(handler)
    built.setLevel(logging.INFO)
    built.propagate = False
    return built


logger = _build_logger()


def log_event(event: str, *, level: int = logging.INFO, **fields: Any) -> None:
    logger.log(level, "ICE_REPORT_COIN_LEDGER_%s", event, extra={"fields": {"event": event, **fields}})


# ---------------------------------------------------------------------------
# Date helpers (Asia/Tokyo, not the container's UTC clock: Cloud Scheduler
# fires at 07:00 JST on the 1st, which is still the previous day in UTC)
# ---------------------------------------------------------------------------


def tokyo_today() -> date:
    return datetime.now(ZoneInfo("Asia/Tokyo")).date()


def previous_month_first(today: date | None = None) -> date:
    today = today or tokyo_today()
    previous_month_last = today.replace(day=1) - timedelta(days=1)
    return previous_month_last.replace(day=1)


def parse_target_month(value: str | None, *, today: date | None = None, required: bool = False) -> date:
    """`YYYY-MM` (the documented form) or `YYYY-MM-01`. Omitted -> previous
    month (scheduled runs) unless `required`. A month that has not finished
    yet in JST can never have complete source data, so it is rejected up
    front rather than left for readiness to (correctly, but less clearly)
    report as not ready."""
    today = today or tokyo_today()
    text = (value or "").strip()
    if not text:
        if required:
            raise CoinLedgerReportError("target_month_required", "target_month (YYYY-MM) is required")
        return previous_month_first(today)

    parsed: date | None = None
    for fmt in ("%Y-%m", "%Y-%m-%d"):
        try:
            parsed = datetime.strptime(text, fmt).date()
            break
        except ValueError:
            continue
    if parsed is None:
        raise CoinLedgerReportError("invalid_target_month", "target_month must be YYYY-MM")
    if parsed.day != 1:
        raise CoinLedgerReportError("invalid_target_month", "target_month must be YYYY-MM")
    if parsed >= today.replace(day=1):
        raise CoinLedgerReportError("invalid_target_month", "target_month must be a completed month")
    return parsed


def output_file_names(target_month: date, generated_date: date) -> dict[str, str]:
    period = f"{target_month:%Y}年{target_month:%m}月期_{generated_date:%y%m%d}"
    return {
        "app": f"【少年ジャンプ＋】消費コイン_{period}.xlsx",
        "web": f"【少年ジャンプ＋】消費コイン_WEB_{period}.xlsx",
        "product": f"話売商品_一覧_{target_month:%y%m}.xlsx",
    }


# ---------------------------------------------------------------------------
# Config accessors (env-overridable, same convention as the other bespoke
# reports)
# ---------------------------------------------------------------------------


def _checked_table(value: str) -> str:
    table = value.strip()
    if table.endswith("_mom"):
        raise CoinLedgerReportError("forbidden_source_table", status_code=500)
    return table


def ledger_table() -> str:
    return _checked_table(os.environ.get("JUMPPLUS_COIN_LEDGER_LEDGER_TABLE", DEFAULT_LEDGER_TABLE))


def content_table() -> str:
    return _checked_table(os.environ.get("JUMPPLUS_COIN_LEDGER_CONTENT_TABLE", DEFAULT_CONTENT_TABLE))


def product_master_table() -> str:
    return _checked_table(os.environ.get("JUMPPLUS_COIN_LEDGER_PRODUCT_MASTER_TABLE", DEFAULT_PRODUCT_MASTER_TABLE))


def output_folder_id() -> str:
    return os.environ.get("JUMPPLUS_COIN_LEDGER_OUTPUT_FOLDER_ID", DEFAULT_OUTPUT_FOLDER_ID).strip()


def template_file_id(file_key: str) -> str:
    for spec in OUTPUT_FILES:
        if spec.file_key == file_key:
            return os.environ.get(spec.template_env, spec.default_template_file_id).strip()
    raise CoinLedgerReportError("unknown_file_key", status_code=500)


def gcs_bucket_name() -> str:
    return (os.environ.get("JUMPPLUS_COIN_LEDGER_BUCKET_NAME") or os.environ.get("BUCKET_NAME") or "").strip()


def gcs_prefix() -> str:
    return os.environ.get("JUMPPLUS_COIN_LEDGER_GCS_PREFIX", DEFAULT_GCS_PREFIX).strip().strip("/")


def delivery_allowed_domains() -> list[str]:
    raw = os.environ.get("JUMPPLUS_COIN_LEDGER_ALLOWED_DOMAINS", "")
    configured = [d.strip().lower() for d in raw.split(",") if d.strip()]
    return configured or list(DEFAULT_ALLOWED_DOMAINS)


# ---------------------------------------------------------------------------
# Readiness
# ---------------------------------------------------------------------------


@dataclass
class ReadinessResult:
    ready: bool
    reasons: list[str] = field(default_factory=list)
    # Safe to log/return: row counts and pass/fail flags only. Coin totals
    # used for reconciliation are kept out of this dict on purpose.
    checks: dict[str, Any] = field(default_factory=dict)

    @property
    def status(self) -> str:
        return "ready" if self.ready else "source_not_ready"


def _query_rows(client: bigquery.Client, query: str, params: list) -> list[dict[str, Any]]:
    job_config = bigquery.QueryJobConfig(query_parameters=params)
    return [dict(row.items()) for row in client.query(query, job_config=job_config).result()]


def _month_param(target_month: date) -> list:
    return [bigquery.ScalarQueryParameter("target_month", "DATE", target_month)]


def evaluate_readiness(
    *,
    ledger_grain: list[dict[str, Any]],
    content_totals: list[dict[str, Any]],
    master_stats: dict[str, Any],
) -> ReadinessResult:
    """Pure evaluation of the aggregate query results (no BigQuery access),
    so every fail-closed rule is unit-testable on its own.

    ledger_grain: one row per (app_pf, coin_type) with row_count, use_coins.
    content_totals: one row per app_id with row_count, total_use_coins.
    master_stats: row_count, distinct_ids, null_ids for content_type='episode'.
    """
    reasons: list[str] = []
    checks: dict[str, Any] = {}

    # ledger: 3 platforms x 6 coin types, exactly one row each.
    grain = {(r.get("app_pf"), r.get("coin_type")): r for r in ledger_grain}
    expected = {(pf, ct) for pf in LEDGER_PLATFORMS for ct in COIN_TYPES}
    missing = sorted(expected - set(grain))
    unexpected = sorted(k for k in grain if k not in expected and k[0] in LEDGER_PLATFORMS)
    duplicated = sorted(k for k, r in grain.items() if k in expected and int(r.get("row_count") or 0) != 1)
    checks["ledger_rows"] = sum(int(r.get("row_count") or 0) for k, r in grain.items() if k in expected)
    checks["ledger_missing_grain_count"] = len(missing)
    checks["ledger_duplicate_grain_count"] = len(duplicated)
    checks["ledger_unexpected_coin_type_count"] = len(unexpected)
    missing_platforms = [pf for pf in LEDGER_PLATFORMS if not any(k[0] == pf for k in grain)]
    if missing_platforms:
        reasons.append("ledger_platform_missing")
    if missing:
        reasons.append("ledger_coin_type_missing")
    if duplicated:
        reasons.append("ledger_duplicate_grain")
    if unexpected:
        reasons.append("ledger_unexpected_coin_type")

    # content: all three app_ids present.
    content_by_app = {int(r["app_id"]): r for r in content_totals if r.get("app_id") is not None}
    for app_id, pf in CONTENT_APP_IDS.items():
        checks[f"content_rows_{pf}"] = int((content_by_app.get(app_id) or {}).get("row_count") or 0)
    if any(app_id not in content_by_app or not content_by_app[app_id].get("row_count") for app_id in CONTENT_APP_IDS):
        reasons.append("content_platform_missing")

    # reconciliation: per platform, content SUM(total_use_coins) ==
    # ledger SUM(use_coins_m).
    reconciled_all = True
    for app_id, pf in CONTENT_APP_IDS.items():
        ledger_use = sum(
            int(r.get("use_coins") or 0) for k, r in grain.items() if k[0] == pf and k in expected
        )
        content_use = int((content_by_app.get(app_id) or {}).get("total_use_coins") or 0)
        ok = app_id in content_by_app and ledger_use == content_use
        checks[f"reconciled_{pf}"] = ok
        reconciled_all = reconciled_all and ok
    if not reconciled_all:
        reasons.append("ledger_content_reconciliation_mismatch")

    # product master: episode rows > 0 and no duplicate / null ids.
    master_rows = int(master_stats.get("row_count") or 0)
    distinct_ids = int(master_stats.get("distinct_ids") or 0)
    null_ids = int(master_stats.get("null_ids") or 0)
    checks["product_master_rows"] = master_rows
    if master_rows <= 0:
        reasons.append("product_master_empty")
    if null_ids or distinct_ids != master_rows - null_ids:
        reasons.append("product_master_duplicate_id")

    return ReadinessResult(ready=not reasons, reasons=reasons, checks=checks)


def check_source_readiness(*, project_id: str, target_month: date, client: bigquery.Client | None = None) -> ReadinessResult:
    client = client or bigquery.Client(project=project_id)
    ledger_grain = _query_rows(
        client,
        f"""
        select app_pf, coin_type, count(*) as row_count, sum(use_coins_m) as use_coins
        from `{ledger_table()}`
        where month_jst = @target_month
        group by app_pf, coin_type
        """,
        _month_param(target_month),
    )
    content_totals = _query_rows(
        client,
        f"""
        select cast(app_id as int64) as app_id, count(*) as row_count, sum(total_use_coins) as total_use_coins
        from `{content_table()}`
        where purchase_date_month_jst = @target_month
            and app_id in (31, 32, 101)
            and purchase_type in unnest(@purchase_types)
        group by app_id
        """,
        _month_param(target_month) + [bigquery.ArrayQueryParameter("purchase_types", "STRING", list(PURCHASE_TYPES))],
    )
    master_rows = _query_rows(
        client,
        f"""
        select count(*) as row_count, count(distinct id) as distinct_ids, countif(id is null) as null_ids
        from `{product_master_table()}`
        where content_type = 'episode'
        """,
        [],
    )
    return evaluate_readiness(
        ledger_grain=ledger_grain,
        content_totals=content_totals,
        master_stats=master_rows[0] if master_rows else {},
    )


# ---------------------------------------------------------------------------
# Source data
# ---------------------------------------------------------------------------

LEDGER_COLUMNS = (
    "month_jst",
    "app_pf",
    "coin_type",
    "coin_type_name",
    "coin_type_num",
    "issue_coins_m",
    "carried_over_coins_m",
    "use_coins_m",
    "cancellation_coins_m",
    "balance_coins_m",
)

CONTENT_COLUMNS = (
    "purchase_date_month_jst",
    "app_id",
    "app_pf",
    "purchase_type",
    "content_id",
    "prefixed_id",
    "v2_content_id_token",
    "name",
    "jdcn",
    "unit_price",
    "download_count",
    "total_use_coins",
    "pay_coins_total",
    "pay_bonus_coins_total",
    "free_ad_coins_total",
    "free_bonus_coins_total",
    "pay_gift_coins_total",
    "reward_video_ad_coin_count",
    "ex_work_name",
    "name_kana",
    "ex_comic_type",
    "ex_sales_start_date",
    "ex_comics_jdcn",
    "ex_episode_package_no",
)

PRODUCT_MASTER_COLUMNS = (
    "prefixed_id",
    "v2_content_id_token",
    "name",
    "work_title",
    "author_name",
    "jdcn",
    "price_in_coin",
    "ex_sales_start_date",
    "ex_comics_jdcn",
    "ex_episode_package_no",
)


@dataclass
class SourceData:
    ledger: list[dict[str, Any]]
    content: list[dict[str, Any]]
    product_master: list[dict[str, Any]]
    comic_type_order_dates: dict[str, dict[str, date | None]] = field(default_factory=dict)


def fetch_ledger_rows(*, client: bigquery.Client, target_month: date) -> list[dict[str, Any]]:
    return _query_rows(
        client,
        f"""
        select {", ".join(LEDGER_COLUMNS)}
        from `{ledger_table()}`
        where month_jst = @target_month
            and app_pf in unnest(@platforms)
            and coin_type in unnest(@coin_types)
        order by app_pf, coin_type_num
        """,
        _month_param(target_month)
        + [
            bigquery.ArrayQueryParameter("platforms", "STRING", list(LEDGER_PLATFORMS)),
            bigquery.ArrayQueryParameter("coin_types", "STRING", list(COIN_TYPES)),
        ],
    )


# Columns that only `_mom` carried (it added them as constant auxiliary
# columns). Fixed to "-" as specified, so the column contract matches the
# existing Excel without reading `_mom`.
MOM_ONLY_FIXED_COLUMNS = {
    "ex_comics_start_date": "-",
    "ex_note": "-",
    "ex_magazine": "-",
    "ex_file_type": "-",
}


def fetch_content_rows(*, client: bigquery.Client, target_month: date) -> list[dict[str, Any]]:
    query = f"""
        select {", ".join(CONTENT_COLUMNS)}
        from `{content_table()}`
        where purchase_date_month_jst = @target_month
            and app_id in (31, 32, 101)
            and purchase_type in unnest(@purchase_types)
        order by app_id, purchase_type, content_id
    """
    job_config = bigquery.QueryJobConfig(
        query_parameters=_month_param(target_month)
        + [bigquery.ArrayQueryParameter("purchase_types", "STRING", list(PURCHASE_TYPES))]
    )
    rows = client.query(query, job_config=job_config).to_dataframe().to_dict("records")
    for row in rows:
        row.update(MOM_ONLY_FIXED_COLUMNS)
    return rows


def fetch_product_master_rows(*, client: bigquery.Client) -> list[dict[str, Any]]:
    query = f"""
        select {", ".join(PRODUCT_MASTER_COLUMNS)}
        from `{product_master_table()}`
        where content_type = 'episode'
        order by id
    """
    return client.query(query).to_dataframe().to_dict("records")


# App サマリの「種別」の対象ラベルは、あくまで当月の購入実績にある
# ex_comic_type(このmodule内 build_summary_rows 側の集計)を正とする --
# raise_master_contents_worksが持つ全履歴の型をそのまま採用すると、
# 当月の購入実績が無い型まで表示してしまい、Golden Masterの大多数の作品
# (当月1種類しか売れていない)で単一値だったものまで誤って複合値に
# なってしまう(2026-08 Golden Runで実際に検証し、この誤りを確認・修正)。
#
# raise_master_contents_worksは「複数の観測ラベルをどの順で連結するか」
# の決定にのみ使う: 各ラベルが最初に episode として現れた日付(無ければ
# 任意の content_type での最初の日付)の昇順。ただし ノベル は常に最後
# (本編/主要なコミック種別に対する副次コンテンツという位置づけ。
# 2026-08 Golden Runで確認したすべての複合値作品は、主要種別が先・
# ノベルが後だった)。この2条件で、2026-08 Golden Masterの複合種別
# 10作品すべての連結順を再現できることを確認済み
# (docs/jumpplus-coin-ledger-report.md 参照)。
def fetch_comic_type_order_dates(*, client: bigquery.Client, work_names: list[str]) -> dict[str, dict[str, date | None]]:
    """{work: {ex_comic_type: sort_date}} -- sort_date is the earliest
    episode-type publish_begin_date_jst for that (work, ex_comic_type), or
    the earliest date of any content_type when that label never appears as
    an episode. Used only to order labels that build_summary_rows already
    observed in this month's purchases; never adds a label the month didn't
    actually have."""
    if not work_names:
        return {}
    query = f"""
        select
          ex_work_name,
          ex_comic_type,
          min(case when content_type = 'episode' then publish_begin_date_jst end) as earliest_episode_date,
          min(publish_begin_date_jst) as earliest_any_date
        from `{product_master_table()}`
        where ex_work_name in unnest(@work_names)
            and ex_comic_type is not null
            and ex_comic_type != ''
        group by ex_work_name, ex_comic_type
    """
    job_config = bigquery.QueryJobConfig(
        query_parameters=[bigquery.ArrayQueryParameter("work_names", "STRING", work_names)]
    )
    rows = [dict(row.items()) for row in client.query(query, job_config=job_config).result()]

    out: dict[str, dict[str, date | None]] = {}
    for row in rows:
        best_date = row["earliest_episode_date"] or row["earliest_any_date"]
        out.setdefault(row["ex_work_name"], {})[row["ex_comic_type"]] = best_date
    return out


def fetch_source_data(*, project_id: str, target_month: date, client: bigquery.Client | None = None) -> SourceData:
    """App/WEB向けのsourceのみ取得する -- `product_master`は空のまま返す
    (`build_workbooks()`がApp/WEB生成完了後、contentを解放してから別途
    取得する)。content(~19万行)とproduct_master(~10万行)を同時に
    メモリ保持し続けることが、Production環境(Linux)でのプロセスRSS
    ピークを大幅に増加させることを実測で確認したため(2026-09、
    docs/jumpplus-coin-ledger-report.md参照)。"""
    client = client or bigquery.Client(project=project_id)
    content = fetch_content_rows(client=client, target_month=target_month)
    work_names = sorted({r["ex_work_name"] for r in content if r.get("ex_work_name")})
    return SourceData(
        ledger=fetch_ledger_rows(client=client, target_month=target_month),
        content=content,
        product_master=[],
        comic_type_order_dates=fetch_comic_type_order_dates(client=client, work_names=work_names),
    )


def content_rows_for(content: list[dict[str, Any]], *, app_id: int, purchase_type: str) -> list[dict[str, Any]]:
    return [
        r
        for r in content
        if r.get("app_id") is not None and int(r["app_id"]) == app_id and r.get("purchase_type") == purchase_type
    ]


# ---------------------------------------------------------------------------
# Output validation (run on every generated file before anything is
# uploaded: a file that fails here never reaches Drive or GCS)
# ---------------------------------------------------------------------------


def validate_xlsx_output(path: str | Path) -> dict[str, Any]:
    from lxml import etree

    with zipfile.ZipFile(path) as zf:
        bad_member = zf.testzip()
        if bad_member is not None:
            raise CoinLedgerReportError("output_zip_corrupt", status_code=500)
        xml_parts = [n for n in zf.namelist() if n.endswith(".xml") or n.endswith(".rels")]
        for name in xml_parts:
            try:
                etree.fromstring(zf.read(name))
            except etree.XMLSyntaxError as exc:
                raise CoinLedgerReportError("output_xml_invalid", status_code=500) from exc

    from openpyxl import load_workbook

    try:
        wb = load_workbook(path, read_only=True)
        sheet_names = list(wb.sheetnames)
        wb.close()
    except Exception as exc:  # openpyxl raises a wide range of types
        raise CoinLedgerReportError("output_readback_failed", status_code=500) from exc
    return {"xml_part_count": len(xml_parts), "sheet_count": len(sheet_names)}


# ---------------------------------------------------------------------------
# Upload helpers
# ---------------------------------------------------------------------------


@dataclass
class GeneratedFile:
    file_key: str
    label: str
    file_name: str
    local_path: Path
    row_counts: dict[str, int]
    drive_file_id: str = ""
    gcs_uri: str = ""


def _upload_to_gcs(local_path: Path, object_name: str) -> str:
    from create_report import upload_to_gcs

    bucket = gcs_bucket_name()
    if not bucket:
        raise CoinLedgerReportError("bucket_not_configured", status_code=500)
    return upload_to_gcs(local_path, bucket, object_name)


def _gcs_object_name(target_month: date, run_stamp: str, file_name: str) -> str:
    return f"{gcs_prefix()}/{target_month:%Y-%m}/{run_stamp}/{file_name}"


def upload_generated_files(
    files: list[GeneratedFile],
    *,
    target_month: date,
    folder_id: str,
    gcs_upload: Callable[[Path, str], str] = _upload_to_gcs,
    drive_upload: Callable[..., dict] | None = None,
    drive_trash: Callable[[str], None] | None = None,
) -> str:
    """Uploads every file to GCS first (private; nothing points at those
    objects until the delivery record is written) and then to the official
    Drive folder. If any Drive upload fails, this run's own earlier Drive
    uploads are trashed (best-effort) so a partial set is never left looking
    like a finished deliverable. Returns the run stamp used in the GCS path."""
    if drive_upload is None or drive_trash is None:
        from drive_io import trash_drive_file, upload_xlsx_to_drive

        drive_upload = drive_upload or upload_xlsx_to_drive
        drive_trash = drive_trash or trash_drive_file

    run_stamp = f"{datetime.now(timezone.utc):%Y%m%dT%H%M%SZ}-{secrets.token_hex(3)}"
    for f in files:
        f.gcs_uri = gcs_upload(f.local_path, _gcs_object_name(target_month, run_stamp, f.file_name))

    uploaded_ids: list[str] = []
    try:
        for f in files:
            uploaded = drive_upload(f.local_path, folder_id=folder_id, file_name=f.file_name, report_type=REPORT_ID)
            f.drive_file_id = str(uploaded.get("id") or "")
            if not f.drive_file_id:
                raise CoinLedgerReportError("drive_upload_failed", status_code=502, retryable=True)
            uploaded_ids.append(f.drive_file_id)
    except Exception:
        for file_id in uploaded_ids:
            try:
                drive_trash(file_id)
            except Exception:
                log_event("DRIVE_CLEANUP_FAILED", level=logging.WARNING)
        raise
    return run_stamp


# ---------------------------------------------------------------------------
# Orchestration
# ---------------------------------------------------------------------------


def build_workbooks(
    *,
    templates: dict[str, Path],
    output_dir: Path,
    target_month: date,
    generated_date: date,
    source: SourceData,
    project_id: str,
    client: bigquery.Client | None = None,
) -> list[GeneratedFile]:
    """Builds and validates all three files locally. Nothing is uploaded
    here, so a failure on any one file leaves no trace outside the temp
    directory.

    `source.product_master` is expected to be empty on entry (see
    `fetch_source_data`): App and WEB only read `source.content` /
    `source.ledger` / `source.comic_type_order_dates`, so this fetches
    product_master (~100k rows) only after App and WEB are built and their
    content/ledger references are dropped, instead of holding content
    (~190k rows) and product_master simultaneously for the whole run. This
    was the dominant driver of Production's OOM (2067 MiB used against a
    2048 MiB limit) -- see docs/jumpplus-coin-ledger-report.md for the
    stage-by-stage RSS measurements that isolated it."""
    import jumpplus_coin_ledger_workbooks as workbooks

    names = output_file_names(target_month, generated_date)
    generated: list[GeneratedFile] = []

    for file_key, builder in (("app", workbooks.build_app_workbook), ("web", workbooks.build_web_workbook)):
        spec = next(s for s in OUTPUT_FILES if s.file_key == file_key)
        output_path = output_dir / names[file_key]
        row_counts = builder(
            template_path=templates[file_key], output_path=output_path, target_month=target_month, source=source
        )
        validate_xlsx_output(output_path)
        generated.append(
            GeneratedFile(
                file_key=file_key, label=spec.label, file_name=names[file_key], local_path=output_path, row_counts=row_counts
            )
        )

    # App/WEB are done: release the large content/ledger structures before
    # fetching product_master, so the two never coexist in memory.
    source.content = []
    source.ledger = []

    client = client or bigquery.Client(project=project_id)
    source.product_master = fetch_product_master_rows(client=client)

    spec = next(s for s in OUTPUT_FILES if s.file_key == "product")
    output_path = output_dir / names["product"]
    row_counts = workbooks.build_product_workbook(
        template_path=templates["product"], output_path=output_path, target_month=target_month, source=source
    )
    validate_xlsx_output(output_path)
    generated.append(
        GeneratedFile(
            file_key="product", label=spec.label, file_name=names["product"], local_path=output_path, row_counts=row_counts
        )
    )
    return generated


def generate_local(
    *,
    project_id: str,
    target_month: date,
    generated_date: date,
    template_dir: Path,
    output_dir: Path,
) -> dict[str, Any]:
    """Golden Run / local verification entry point: BigQuery read + local
    workbook build only. Never touches Drive, GCS or Firestore."""
    readiness = check_source_readiness(project_id=project_id, target_month=target_month)
    if not readiness.ready:
        return {"status": "source_not_ready", "reasons": readiness.reasons, "checks": readiness.checks}
    source = fetch_source_data(project_id=project_id, target_month=target_month)
    templates = {key: template_dir / f"{key}.xlsx" for key in OUTPUT_FILE_KEYS}
    output_dir.mkdir(parents=True, exist_ok=True)
    files = build_workbooks(
        templates=templates,
        output_dir=output_dir,
        target_month=target_month,
        generated_date=generated_date,
        source=source,
        project_id=project_id,
    )
    return {
        "status": "ok",
        "checks": readiness.checks,
        "files": [{"file_key": f.file_key, "path": str(f.local_path), "row_counts": f.row_counts} for f in files],
    }


def generate_coin_ledger_report(
    *,
    project_id: str,
    target_month: date,
    generated_date: date,
    trigger: str,
    deliver: Callable[[list[GeneratedFile]], dict[str, Any]] | None = None,
) -> dict[str, Any]:
    """readiness -> BigQuery -> 3 workbooks (local, validated) -> GCS -> Drive
    -> `deliver` (writes the multi-file delivery version). `deliver` runs
    only after all three files exist in both GCS and Drive, and is the only
    step that makes anything reachable from a delivery URL."""
    if not project_id:
        raise CoinLedgerReportError("project_required", status_code=500)
    started = time.monotonic()
    log_event("STARTED", target_month=target_month.isoformat(), trigger=trigger)

    readiness = check_source_readiness(project_id=project_id, target_month=target_month)
    log_event(
        "READINESS",
        level=logging.INFO if readiness.ready else logging.WARNING,
        target_month=target_month.isoformat(),
        trigger=trigger,
        readiness=readiness.status,
        reasons=readiness.reasons,
        **readiness.checks,
    )
    if not readiness.ready:
        raise CoinLedgerReportError(
            "source_not_ready",
            status_code=503,
            retryable=True,
            reasons=readiness.reasons,
            checks=readiness.checks,
        )

    source = fetch_source_data(project_id=project_id, target_month=target_month)

    from drive_io import download_drive_file

    folder_id = output_folder_id()
    with tempfile.TemporaryDirectory(prefix="coin-ledger-") as tmp_dir:
        tmp_root = Path(tmp_dir)
        templates: dict[str, Path] = {}
        for key in OUTPUT_FILE_KEYS:
            file_id = template_file_id(key)
            if not file_id:
                raise CoinLedgerReportError("template_not_configured", status_code=500)
            templates[key] = download_drive_file(file_id, tmp_root / "templates" / f"{key}.xlsx")

        output_dir = tmp_root / "output"
        output_dir.mkdir()
        files = build_workbooks(
            templates=templates,
            output_dir=output_dir,
            target_month=target_month,
            generated_date=generated_date,
            source=source,
            project_id=project_id,
        )
        log_event(
            "WORKBOOKS_BUILT",
            target_month=target_month.isoformat(),
            trigger=trigger,
            row_counts={f.file_key: f.row_counts for f in files},
        )

        upload_generated_files(files, target_month=target_month, folder_id=folder_id)
        log_event(
            "DRIVE_UPLOAD_COMPLETED",
            target_month=target_month.isoformat(),
            trigger=trigger,
            file_count=len(files),
        )

        delivery_result = deliver(files) if deliver else None

    elapsed = round(time.monotonic() - started, 1)
    log_event(
        "COMPLETED",
        target_month=target_month.isoformat(),
        trigger=trigger,
        delivery_id=(delivery_result or {}).get("delivery_id", ""),
        delivery_version=(delivery_result or {}).get("version"),
        elapsed_seconds=elapsed,
    )
    return {
        "status": "ok",
        "report": REPORT_ID,
        "report_name": REPORT_NAME,
        "target_month": target_month.strftime("%Y-%m"),
        "generated_date": generated_date.isoformat(),
        "trigger": trigger,
        "readiness": {"status": readiness.status, **readiness.checks},
        "files": [
            {
                "file_key": f.file_key,
                "label": f.label,
                "file_name": f.file_name,
                "drive_file_id": f.drive_file_id,
                "has_gcs_object": bool(f.gcs_uri),
                "row_counts": f.row_counts,
            }
            for f in files
        ],
        "delivery": delivery_result,
        "elapsed_seconds": elapsed,
    }


def _main(argv: list[str] | None = None) -> int:
    import argparse

    parser = argparse.ArgumentParser(description="Local Golden Run: BigQuery read + local XLSX build only.")
    parser.add_argument("--target-month", required=True)
    parser.add_argument("--generated-date", required=True, help="YYYY-MM-DD used in output file names")
    parser.add_argument("--template-dir", required=True, help="directory holding app.xlsx / web.xlsx / product.xlsx")
    parser.add_argument("--output-dir", required=True)
    parser.add_argument("--project-id", default=os.environ.get("BIGQUERY_PROJECT_ID", "jumpplus-4a5f4"))
    args = parser.parse_args(argv)

    generated_date = datetime.strptime(args.generated_date, "%Y-%m-%d").date()
    target_month = parse_target_month(args.target_month, today=tokyo_today(), required=True)
    result = generate_local(
        project_id=args.project_id,
        target_month=target_month,
        generated_date=generated_date,
        template_dir=Path(args.template_dir),
        output_dir=Path(args.output_dir),
    )
    print(json.dumps(result, ensure_ascii=False, indent=2, default=str))
    return 0 if result["status"] == "ok" else 2


if __name__ == "__main__":
    raise SystemExit(_main())
