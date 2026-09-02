from __future__ import annotations

import copy
import os
import re
import tempfile
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any
from zoneinfo import ZoneInfo

from google.cloud import bigquery
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter, range_boundaries
from openpyxl.worksheet.worksheet import Worksheet

PLUS_REPORT_ID = "plus-browser-point-sales"
PLUS_REPORT_NAME = "PLUS_ブラウザ版_ポイント売上（月次Excel反映）"

DEFAULT_PLUS_TEMPLATE_FILE_ID = "1y2axBdZVfkeeJ4GiA76k8jqra4uAPiFf"
DEFAULT_PLUS_OUTPUT_FOLDER_ID = "1FjFRdkQZz6yI5sRe00RSBiG3pGEwTNq_"
DEFAULT_PLUS_SOURCE_TABLE = "jumpplus-4a5f4.dataset_process_tables.daily_sbps_order_combined"
DEFAULT_PLUS_PRODUCT_TABLE = "jumpplus-4a5f4.dataset_exdata_tables.sbps_product"
DEFAULT_PLUS_PAYMENT_CLASS_TABLE = "jumpplus-4a5f4.dataset_exdata_tables.sbps_payment_class"

BY_PAYMENT_SHEET = "合計_決済別"
BY_PAYMENT_TABLE = "合計_決済別"
BY_PAYMENT_PRODUCT_SHEET = "合計_決済-商品別"
BY_PAYMENT_PRODUCT_TABLE = "合計_決済_商品別"

# Official price -> Excel point-label mapping (confirmed business spec, not a
# heuristic). `price` is the purchase amount in yen and the aggregation key
# throughout this module; BigQuery's `product_name` (e.g. "300pt") is a
# separate, purely informational label local to `daily_sbps_order_combined` /
# `sbps_product` and is never used for the Excel mapping. The Excel template's
# `合計_決済-商品別` columns use the bonus-inclusive "granted points" label for
# each price tier (e.g. buying the 300-yen tier grants 310pt, i.e. a 10pt
# bonus) -- that is the value on the right below, and it is the only source of
# truth for the column header text.
#
# This map is intentionally static. When a new price tier is introduced, it
# must be added here (and to the Excel template) explicitly by an engineer --
# `build_report_data` fails closed with `unexpected_product_price` for any
# `price` it does not recognize, rather than inferring a new label.
PRICE_TO_POINT_LABEL: dict[int, str] = {
    100: "100pt",
    300: "310pt",
    400: "410pt",
    500: "520pt",
    600: "630pt",
    1000: "1050pt",
    2000: "2150pt",
    3000: "3250pt",
    5000: "5450pt",
    9800: "10800pt",
}

PAYMENT_CLASS_PREFIX_RE = re.compile(r"^(\d+)_(.+)$")

PLUS_SQL = """
with actual as (
    select
        price
        , payment_class
        , sum(sum_price) as sum_price
    from
        `{source_table}`
    where
        month = @target_month
    group by all
)

, product_master as (
    select distinct
        price
        , product_name
    from
        `{product_table}`
)

, payment_master as (
    select distinct
        payment_class
    from
        `{payment_class_table}`
    where
        payment_class is not null
)

, zero_fill as (
    select
        product_master.price
        , payment_master.payment_class
        , 0 as sum_price
    from
        product_master
    cross join
        payment_master
)

, combined as (
    select price, payment_class, sum_price from actual
    union all
    select price, payment_class, sum_price from zero_fill
)

select
    price
    , payment_class
    , sum(sum_price) as sum_price
from
    combined
group by all
"""


class PlusPointSalesReportError(Exception):
    def __init__(self, code: str, message: str | None = None, *, status_code: int = 400, **details: Any) -> None:
        super().__init__(message or code)
        self.code = code
        self.status_code = status_code
        self.details = details


def tokyo_today() -> date:
    return datetime.now(ZoneInfo("Asia/Tokyo")).date()


def previous_month_first(today: date | None = None) -> date:
    today = today or tokyo_today()
    first_this_month = today.replace(day=1)
    previous_month_last = first_this_month - timedelta(days=1)
    return previous_month_last.replace(day=1)


def parse_target_month(value: str | None = None, *, today: date | None = None) -> date:
    if not value:
        return previous_month_first(today)
    try:
        parsed = datetime.strptime(value, "%Y-%m-%d").date()
    except ValueError as exc:
        raise PlusPointSalesReportError("invalid_target_month", "target_month must be YYYY-MM-DD") from exc
    if parsed.day != 1:
        raise PlusPointSalesReportError("invalid_target_month", "target_month must be the first day of month")
    return parsed


def output_file_name(target_month: date) -> str:
    return f"J+ブラウザ版_ポイント売上_{target_month:%y年%m月}分.xlsx"


def source_table() -> str:
    return os.environ.get("PLUS_POINT_SALES_SOURCE_TABLE", DEFAULT_PLUS_SOURCE_TABLE)


def product_table() -> str:
    return os.environ.get("PLUS_POINT_SALES_PRODUCT_TABLE", DEFAULT_PLUS_PRODUCT_TABLE)


def payment_class_table() -> str:
    return os.environ.get("PLUS_POINT_SALES_PAYMENT_CLASS_TABLE", DEFAULT_PLUS_PAYMENT_CLASS_TABLE)


def _payment_class_sort_key(payment_class: str) -> tuple[int, str]:
    match = PAYMENT_CLASS_PREFIX_RE.match(payment_class)
    if not match:
        raise PlusPointSalesReportError(
            "unparseable_payment_class",
            "payment_class does not match the expected '<number>_<label>' format",
            payment_class=payment_class,
        )
    return int(match.group(1)), payment_class


def payment_class_display_label(payment_class: str) -> str:
    match = PAYMENT_CLASS_PREFIX_RE.match(payment_class)
    if not match:
        raise PlusPointSalesReportError(
            "unparseable_payment_class",
            "payment_class does not match the expected '<number>_<label>' format",
            payment_class=payment_class,
        )
    return match.group(2)


def run_plus_point_sales_query(
    *,
    project_id: str,
    target_month: date,
    source: str | None = None,
    products: str | None = None,
    payment_classes: str | None = None,
) -> list[dict[str, Any]]:
    client = bigquery.Client(project=project_id)
    query = PLUS_SQL.format(
        source_table=source or source_table(),
        product_table=products or product_table(),
        payment_class_table=payment_classes or payment_class_table(),
    )
    job_config = bigquery.QueryJobConfig(
        query_parameters=[
            bigquery.ScalarQueryParameter("target_month", "DATE", target_month),
        ]
    )
    frame = client.query(query, job_config=job_config).to_dataframe()
    return frame.to_dict("records")


def build_report_data(records: list[dict[str, Any]]) -> dict[str, Any]:
    """Validate raw query rows and pivot them into the two report tables.

    Fails closed (raises PlusPointSalesReportError) on any payment_class or
    price that this module does not already understand, instead of silently
    dropping or misreporting yen amounts.
    """
    by_payment_product: dict[str, dict[int, int]] = {}
    by_payment_total: dict[str, int] = {}
    payment_classes: set[str] = set()

    for row in records:
        payment_class = row.get("payment_class")
        price_raw = row.get("price")
        sum_price = int(row.get("sum_price") or 0)

        if payment_class is None:
            if sum_price != 0:
                raise PlusPointSalesReportError(
                    "unexpected_payment_class",
                    "sales row has no mapped payment_class in sbps_payment_class",
                    price=price_raw,
                    sum_price=sum_price,
                )
            continue

        if price_raw is None:
            raise PlusPointSalesReportError(
                "unexpected_product_price",
                "sales row has a mapped payment_class but a null price",
                payment_class=payment_class,
            )
        price = int(price_raw)
        if price not in PRICE_TO_POINT_LABEL:
            # Fails closed unconditionally (even for a sum_price == 0 zero-fill
            # row) so a newly introduced sbps_product price tier is caught the
            # first month it appears in the master table, before it could ever
            # be silently mapped to the wrong Excel column. The fix is to add
            # the new price -> point-label pair to PRICE_TO_POINT_LABEL (and
            # the Excel template's column) explicitly, never to guess it here.
            raise PlusPointSalesReportError(
                "unexpected_product_price",
                "sbps_product price has no known point-label mapping in PRICE_TO_POINT_LABEL",
                price=price,
                payment_class=payment_class,
            )

        # Validates the numeric-prefix format up front so a malformed
        # payment_class is reported clearly rather than sorted wrong later.
        _payment_class_sort_key(payment_class)

        payment_classes.add(payment_class)
        by_payment_product.setdefault(payment_class, {})
        by_payment_product[payment_class][price] = by_payment_product[payment_class].get(price, 0) + sum_price
        by_payment_total[payment_class] = by_payment_total.get(payment_class, 0) + sum_price

    ordered_payment_classes = sorted(payment_classes, key=_payment_class_sort_key)
    ordered_prices = sorted(PRICE_TO_POINT_LABEL.keys())

    grand_total_by_payment = sum(by_payment_total.values())
    grand_total_by_product = sum(
        by_payment_product.get(pc, {}).get(price, 0)
        for pc in ordered_payment_classes
        for price in ordered_prices
    )
    if grand_total_by_payment != grand_total_by_product:
        raise PlusPointSalesReportError(
            "grand_total_mismatch",
            "sum by payment_class does not match sum by price x payment_class",
            grand_total_by_payment=grand_total_by_payment,
            grand_total_by_product=grand_total_by_product,
        )

    return {
        "payment_classes": ordered_payment_classes,
        "prices": ordered_prices,
        "by_payment_total": by_payment_total,
        "by_payment_product": by_payment_product,
        "grand_total": grand_total_by_payment,
    }


def _extend_table_rows(ws: Worksheet, table_name: str, target_data_row_count: int) -> None:
    """Grow an Excel Table's data rows in place, ahead of its totals row.

    The template's SUBTOTAL formulas use whole-column structured references
    (e.g. `SUBTOTAL(109,合計_決済別[売上])`), so they keep summing correctly
    once the table's `ref` covers the new rows -- no formula text changes
    needed. New rows copy cell style/number-format from the row directly
    above the (old) totals row so banding/number formatting is preserved.
    """
    table = ws.tables[table_name]
    min_col, min_row, max_col, max_row = range_boundaries(table.ref)
    header_row = min_row
    old_totals_row = max_row
    current_data_rows = old_totals_row - header_row - 1
    extra = target_data_row_count - current_data_rows
    if extra < 0:
        raise PlusPointSalesReportError(
            "template_row_shrink_unsupported",
            "target payment_class count is smaller than the template's existing rows",
            table_name=table_name,
        )
    if extra == 0:
        return

    style_source_row = old_totals_row - 1
    ws.insert_rows(old_totals_row, amount=extra)

    for offset in range(extra):
        new_row = old_totals_row + offset
        for col in range(min_col, max_col + 1):
            src = ws.cell(row=style_source_row, column=col)
            dst = ws.cell(row=new_row, column=col)
            if src.has_style:
                dst._style = copy.copy(src._style)
            dst.number_format = src.number_format
        if style_source_row in ws.row_dimensions and ws.row_dimensions[style_source_row].height:
            ws.row_dimensions[new_row].height = ws.row_dimensions[style_source_row].height

    new_totals_row = old_totals_row + extra
    table.ref = f"{get_column_letter(min_col)}{header_row}:{get_column_letter(max_col)}{new_totals_row}"
    if table.autoFilter is not None:
        table.autoFilter.ref = f"{get_column_letter(min_col)}{header_row}:{get_column_letter(max_col)}{new_totals_row - 1}"


def write_by_payment_sheet(ws: Worksheet, data: dict[str, Any]) -> None:
    payment_classes = data["payment_classes"]
    _extend_table_rows(ws, BY_PAYMENT_TABLE, len(payment_classes))
    header_row = 1
    for offset, payment_class in enumerate(payment_classes):
        row = header_row + 1 + offset
        ws.cell(row=row, column=1).value = payment_class_display_label(payment_class)
        ws.cell(row=row, column=2).value = int(data["by_payment_total"].get(payment_class, 0))


def write_by_payment_product_sheet(ws: Worksheet, data: dict[str, Any]) -> None:
    payment_classes = data["payment_classes"]
    prices = data["prices"]
    _extend_table_rows(ws, BY_PAYMENT_PRODUCT_TABLE, len(payment_classes))
    header_row = 1
    total_col = len(prices) + 2  # column A = label, B..K = products, L = TOTAL
    row_total_formula = (
        f"=SUM({BY_PAYMENT_PRODUCT_TABLE}[@[{PRICE_TO_POINT_LABEL[prices[0]]}]:[{PRICE_TO_POINT_LABEL[prices[-1]]}]])"
    )

    for offset, payment_class in enumerate(payment_classes):
        row = header_row + 1 + offset
        ws.cell(row=row, column=1).value = payment_class_display_label(payment_class)
        product_values = data["by_payment_product"].get(payment_class, {})
        for col_offset, price in enumerate(prices):
            ws.cell(row=row, column=2 + col_offset).value = int(product_values.get(price, 0))
        ws.cell(row=row, column=total_col).value = row_total_formula


def create_plus_point_sales_workbook(
    *,
    template_path: str | Path,
    output_path: str | Path,
    data: dict[str, Any],
) -> None:
    workbook = load_workbook(template_path)
    if BY_PAYMENT_SHEET not in workbook.sheetnames:
        raise PlusPointSalesReportError("by_payment_sheet_not_found", f"{BY_PAYMENT_SHEET} sheet not found")
    if BY_PAYMENT_PRODUCT_SHEET not in workbook.sheetnames:
        raise PlusPointSalesReportError(
            "by_payment_product_sheet_not_found", f"{BY_PAYMENT_PRODUCT_SHEET} sheet not found"
        )

    write_by_payment_sheet(workbook[BY_PAYMENT_SHEET], data)
    write_by_payment_product_sheet(workbook[BY_PAYMENT_PRODUCT_SHEET], data)

    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)


def generate_plus_point_sales_report(
    *,
    project_id: str,
    target_month_text: str | None = None,
    today: date | None = None,
    template_file_id: str | None = None,
    output_folder_id: str | None = None,
) -> dict[str, Any]:
    if not project_id:
        raise PlusPointSalesReportError("project_required", "BIGQUERY_PROJECT_ID or PROJECT_ID is required")

    today = today or tokyo_today()
    target_month = parse_target_month(target_month_text, today=today)
    template_id = template_file_id or os.environ.get(
        "PLUS_POINT_SALES_TEMPLATE_FILE_ID", DEFAULT_PLUS_TEMPLATE_FILE_ID
    )
    folder_id = output_folder_id or os.environ.get(
        "PLUS_POINT_SALES_OUTPUT_FOLDER_ID", DEFAULT_PLUS_OUTPUT_FOLDER_ID
    )
    file_name = output_file_name(target_month)

    records = run_plus_point_sales_query(project_id=project_id, target_month=target_month)
    data = build_report_data(records)

    from drive_io import download_drive_file, upload_xlsx_to_drive

    with tempfile.TemporaryDirectory(prefix="plus-point-sales-report-") as tmp_dir:
        tmp_root = Path(tmp_dir)
        template_path = tmp_root / "template.xlsx"
        output_path = tmp_root / file_name
        download_drive_file(template_id, template_path)
        create_plus_point_sales_workbook(
            template_path=template_path,
            output_path=output_path,
            data=data,
        )
        uploaded = upload_xlsx_to_drive(
            output_path,
            folder_id=folder_id,
            file_name=file_name,
        )

    return {
        "status": "ok",
        "report": PLUS_REPORT_ID,
        "report_name": PLUS_REPORT_NAME,
        "target_month": target_month.isoformat(),
        "generated_date": today.isoformat(),
        "file_id": uploaded.get("id", ""),
        "file_name": uploaded.get("name") or file_name,
        "webViewLink": uploaded.get("webViewLink", ""),
        "payment_class_count": len(data["payment_classes"]),
        "grand_total": data["grand_total"],
    }
