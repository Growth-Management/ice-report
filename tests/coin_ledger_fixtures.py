"""Dummy data and synthetic templates for the jumpplus-coin-ledger tests.

Everything here is invented (no production values). The synthetic
templates mirror the official templates' structure as analysed on
2026-09-28 (sheet names, Table names/headers, fixed 出納 layout, totals
rows, freeze panes, zoom) closely enough to exercise the writers; the real
templates are only used by the opt-in test that reads
JUMPPLUS_COIN_LEDGER_TEMPLATE_DIR, never committed to the repository.
"""

from __future__ import annotations

import random
from datetime import date
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.worksheet.table import Table, TableStyleInfo

import jumpplus_coin_ledger_report as report
import jumpplus_coin_ledger_workbooks as workbooks

APP_IDS = {31: "iOS", 32: "And", 101: "Web"}


def dummy_source(*, seed: int = 7, contents_per_app: int = 40, works: int = 6, master_rows: int = 25) -> report.SourceData:
    rng = random.Random(seed)
    content = []
    for app_id, pf in APP_IDS.items():
        for i in range(contents_per_app):
            purchase_type = "episode" if i % 3 else "book"
            parts = [rng.randint(0, 50) for _ in range(6)]
            content.append(
                {
                    "purchase_date_month_jst": date(2026, 8, 1),
                    "app_id": Decimal(app_id),
                    "app_pf": pf,
                    "purchase_type": purchase_type,
                    "content_id": Decimal(1000 + i),
                    "prefixed_id": f"{purchase_type}-{1000 + i}",
                    "v2_content_id_token": f"tok{1000 + i:05d}",
                    "name": f"ダミー作品{i % works} 第{i}話",
                    "jdcn": f"J{i:07d}",
                    "unit_price": Decimal(rng.choice([30, 50, 60])),
                    "download_count": Decimal(rng.randint(1, 20)),
                    "total_use_coins": Decimal(sum(parts)),
                    "pay_coins_total": Decimal(parts[0]),
                    "pay_bonus_coins_total": Decimal(parts[1]),
                    "free_ad_coins_total": Decimal(parts[2]),
                    "free_bonus_coins_total": Decimal(parts[3]),
                    "pay_gift_coins_total": Decimal(parts[4]),
                    "reward_video_ad_coin_count": Decimal(parts[5]),
                    "ex_work_name": f"ダミー作品{i % works}",
                    "name_kana": "ダミー",
                    "ex_comic_type": "オリジナル" if i % works else "コミックス",
                    "ex_sales_start_date": date(2025, 1 + i % 12, 1),
                    "ex_comics_jdcn": None if i % 4 else f"C{i:06d}",
                    "ex_episode_package_no": None if i % 4 else i % 9,
                    **report.MOM_ONLY_FIXED_COLUMNS,
                }
            )

    ledger = []
    for pf in report.LEDGER_PLATFORMS:
        use_total = int(sum(r["total_use_coins"] for r in content if r["app_pf"] == pf))
        for n, coin_type in enumerate(report.COIN_TYPES, start=1):
            use = use_total // 6 + (use_total % 6 if n == 6 else 0)
            issue, carried, cancel = 1000 * n, 500 * n, -n
            ledger.append(
                {
                    "month_jst": date(2026, 8, 1),
                    "app_pf": pf,
                    "coin_type": coin_type,
                    "coin_type_name": coin_type,
                    "coin_type_num": n,
                    "issue_coins_m": Decimal(issue),
                    "carried_over_coins_m": Decimal(carried),
                    "use_coins_m": Decimal(use),
                    "cancellation_coins_m": Decimal(cancel),
                    "balance_coins_m": Decimal(carried + issue - use + cancel),
                }
            )

    master = [
        {
            "prefixed_id": f"episode-{i}",
            "v2_content_id_token": f"tok{i:05d}",
            "name": f"ダミー話{i}",
            "work_title": f"ダミー作品{i % 5}",
            "author_name": f"著者{i % 3}",
            "jdcn": f"J{i:07d}",
            "price_in_coin": 50,
            "ex_sales_start_date": date(2024, 1 + i % 12, 1),
            "ex_comics_jdcn": None,
            "ex_episode_package_no": None,
        }
        for i in range(master_rows)
    ]
    return report.SourceData(ledger=ledger, content=content, product_master=master)


# ---------------------------------------------------------------------------
# Synthetic templates
# ---------------------------------------------------------------------------


def _ledger_block(ws, block: workbooks.LedgerBlock, labels, *, type_header: bool) -> None:
    ws[block.title_ref] = block.title
    h = block.header_row
    if type_header:
        ws[f"A{h}"] = "コイン種別"
    for col, text in {**workbooks.LEDGER_HEADER_TEXTS, "G": "当月消費率", "H": "繰越込消費率"}.items():
        ws[f"{col}{h}"] = text
    first = block.first_coin_row
    for offset, label in enumerate(labels):
        r = first + offset
        ws[f"A{r}"] = label
        for col in "BCDE":
            ws[f"{col}{r}"].number_format = "#,##0"
        ws[f"F{r}"] = f"=C{r}+B{r}-D{r}+E{r}"
        ws[f"G{r}"] = f'=IFERROR(D{r}/B{r},"-")'
        ws[f"H{r}"] = f'=IFERROR(D{r}/SUM(B{r}:C{r}),"-")'
    total = first + len(labels)
    ws[f"A{total}"] = "総計"
    for col in "BCDEF":
        ws[f"{col}{total}"] = f"=SUM({col}{first}:{col}{total - 1})"


def _table_sheet(wb, title, table_name, headers, *, blank_rows=1, totals=False, freeze=None, zoom=85):
    ws = wb.create_sheet(title)
    ws["A1"] = title
    for c, header in enumerate(headers, start=1):
        ws.cell(row=3, column=c, value=header)
        for r in range(4, 4 + blank_rows):
            ws.cell(row=r, column=c).number_format = "#,##0"
    last_data = 3 + blank_rows
    last_col = ws.cell(row=3, column=len(headers)).column_letter
    ref = f"A3:{last_col}{last_data + (1 if totals else 0)}"
    table = Table(displayName=table_name, ref=ref)
    table.tableStyleInfo = TableStyleInfo(name="TableStyleLight8", showRowStripes=True)
    if totals:
        table.totalsRowCount = 1
        ws.cell(row=last_data + 1, column=1, value="総計")
        for c, header in enumerate(headers[1:14], start=2):
            ws.cell(row=last_data + 1, column=c, value=f"=SUBTOTAL(109,{table_name}[{header}])")
    ws.add_table(table)
    if totals:
        # openpyxl only generates tableColumns at save time; build them now
        # so the totals-row metadata can be set in the same shape the real
        # template has.
        table._initialise_columns()
        for column, header in zip(table.tableColumns, headers):
            column.name = header
        table.tableColumns[0].totalsRowLabel = "総計"
        for column in table.tableColumns[1:14]:
            column.totalsRowFunction = "sum"
    if freeze:
        ws.freeze_panes = freeze
    ws.sheet_view.zoomScale = zoom
    return ws


def build_app_template(path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "出納"
    ws.sheet_view.zoomScale = 85
    for block in workbooks.APP_LEDGER_BLOCKS:
        _ledger_block(ws, block, workbooks.APP_COIN_LABELS, type_header=True)
    for sheet_name, table_name, with_type in (
        ("サマリ", "サマリ_Total", True),
        ("サマリ (Apple)", "サマリ_Apple", False),
        ("サマリ (Google)", "サマリ_Google", False),
    ):
        _table_sheet(wb, sheet_name, table_name, workbooks.summary_headers("コイン", with_type=with_type), blank_rows=30, totals=True)
    for (sheet_name, _, purchase_type), table_name in zip(
        workbooks.APP_DETAIL_SHEETS, ("話明細_Apple", "話明細_Google", "巻明細_Apple", "巻明細_Google")
    ):
        _table_sheet(wb, sheet_name, table_name, workbooks.detail_headers("コイン", purchase_type), freeze="E4", zoom=80)
    wb.save(path)
    return path


def build_web_template(path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "出納"
    for block in workbooks.WEB_LEDGER_BLOCKS:
        _ledger_block(ws, block, workbooks.WEB_COIN_LABELS, type_header=False)
    _table_sheet(wb, "サマリ", "サマリ_Total", workbooks.summary_headers("ポイント", with_type=True), totals=True)
    for (sheet_name, _, purchase_type), table_name in zip(workbooks.WEB_DETAIL_SHEETS, ("話明細_Apple", "巻明細_Apple")):
        _table_sheet(wb, sheet_name, table_name, workbooks.detail_headers("ポイント", purchase_type), freeze="E4", zoom=80)
    wb.save(path)
    return path


def build_product_template(path: Path) -> Path:
    """Mirrors the official 話売商品_一覧_yymm.xlsx: one sheet "file", Table
    "テーブル1" with its header on row 1 (A1:J2), frozen at A2, zoom 85."""
    wb = Workbook()
    ws = wb.active
    ws.title = "file"
    headers = tuple(h for h, _ in workbooks.PRODUCT_HEADERS)
    for c, header in enumerate(headers, start=1):
        ws.cell(row=1, column=c, value=header)
    table = Table(displayName="テーブル1", ref="A1:J2")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleLight21", showRowStripes=True)
    ws.add_table(table)
    ws.freeze_panes = "A2"
    ws.sheet_view.zoomScale = 85
    wb.save(path)
    return path
