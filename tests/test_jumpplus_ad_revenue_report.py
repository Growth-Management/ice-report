import importlib
import io
import sys
import tempfile
import unittest
from datetime import date
from pathlib import Path
from unittest import mock

from openpyxl import Workbook
from openpyxl.worksheet.table import Table, TableColumn

import jumpplus_ad_revenue_report as report

# Sibling test modules stub out "openpyxl" and/or "google"/"google.cloud" with
# incomplete fakes (see test_plus_browser_point_sales_report.py's comment on
# the identical openpyxl issue). Some do this inside a test method (pollution
# only visible to tests that run afterward); test_admin_iap_auth.py's
# `_install_import_stubs()` does it at module *import* time via
# `sys.modules.setdefault(...)`, which can land before this file is even
# collected depending on pytest's file order -- confirmed empirically: by the
# time this file's own `import jumpplus_ad_revenue_report` line above runs,
# `report.bigquery` can already be that stub's empty `SimpleNamespace()`.
# A plain "capture what's in sys.modules right now" is therefore not reliable
# for the google.* family; evict any google.* entries and force a genuine
# reimport (the real google-cloud-bigquery package is actually installed) to
# get a trustworthy reference, then put back whatever was there so later test
# files still see what they expect.
def _reimport_real(root: str, target: str):
    """Returns (module, fresh_tree): `module` is the freshly, genuinely
    imported `target`; `fresh_tree` is every real `root`/`root.*` module this
    reimport populated in sys.modules. bigquery's own classes do further
    *internal* lazy cross-submodule imports at call time (e.g.
    `Table.__init__` and `StructQueryParameter.from_api_repr` each do their
    own `from google.cloud.bigquery import ...`), so it is not enough to fix
    up this test file's own `report.bigquery` name -- the whole real
    "google.*" tree must be reinstated in sys.modules for the duration of any
    test that exercises real bigquery object construction, not just captured
    once as a local reference.

    Only call this when the currently-cached module is demonstrably broken
    (see _ensure_real_bigquery below). Evicting and reimporting an
    already-loaded, healthy "google.*" tree is not safe in general -- grpc/
    protobuf's C-extension state does not reliably support being torn down
    and reinitialized mid-process (confirmed empirically: doing this
    unconditionally, in a process where nothing had polluted sys.modules yet,
    crashed the interpreter).
    """
    saved = {n: m for n, m in sys.modules.items() if n == root or n.startswith(root + ".")}
    for n in saved:
        del sys.modules[n]
    try:
        module = importlib.import_module(target)
        fresh_tree = {n: m for n, m in sys.modules.items() if n == root or n.startswith(root + ".")}
    finally:
        for n in list(sys.modules):
            if n == root or n.startswith(root + "."):
                del sys.modules[n]
        sys.modules.update(saved)
    return module, fresh_tree


def _ensure_real_bigquery():
    """`jumpplus_ad_revenue_report.bigquery` can already be a broken stub by
    the time this test file is even collected (test_admin_iap_auth.py installs
    one at *import* time via `sys.modules.setdefault(...)`, which can land
    before this file is collected depending on pytest's file order) -- or it
    can still be the genuine module if nothing has polluted things yet. Only
    take the _reimport_real path when it is demonstrably broken (real
    bigquery always has `SchemaField`; every stub seen in this codebase's
    tests does not), so a healthy environment is never touched.
    """
    if hasattr(report.bigquery, "SchemaField"):
        return report.bigquery, {}
    return _reimport_real("google", "google.cloud.bigquery")


_REAL_BIGQUERY, _REAL_GOOGLE_MODULES = _ensure_real_bigquery()
# openpyxl pollution (unlike google.*) has only ever been observed to happen
# during test *execution*, so a straightforward capture at this file's import
# time is sufficient -- mirrors test_plus_browser_point_sales_report.py.
_REAL_OPENPYXL_MODULES = {
    name: mod for name, mod in sys.modules.items() if name == "openpyxl" or name.startswith("openpyxl.")
}


def _restore_real_dependencies() -> None:
    sys.modules.update(_REAL_OPENPYXL_MODULES)
    sys.modules.update(_REAL_GOOGLE_MODULES)
    report.bigquery = _REAL_BIGQUERY


class _RestoringTestCase(unittest.TestCase):
    def setUp(self):
        _restore_real_dependencies()


def _build_summary_sheet(wb, *, with_multiple_total_blocks=False):
    ws = wb.create_sheet(report.SUMMARY_SHEET)
    ws["A1"] = "動画リワード広告売上"
    ws["B3"] = "区分"
    ws["C3"] = "広告売上"
    ws["B4"] = "総計"
    ws["C4"] = None
    if with_multiple_total_blocks:
        ws["B6"] = "区分"
        ws["C6"] = "コイン消費数"
        ws["B7"] = "iOS"
        ws["C7"] = "#REF!"
        ws["B8"] = "Android"
        ws["C8"] = "#REF!"
        ws["B9"] = "総計"
        ws["C9"] = "#REF!"
        ws["B11"] = "区分"
        ws["C11"] = "還元額"
        ws["B12"] = "総計"
        ws["C12"] = "-"
    return ws


def _build_detail_sheet(wb, sheet_name, headers, data_row_count, *, extra_formula_headers=None):
    ws = wb.create_sheet(sheet_name)
    all_headers = list(headers) + list(extra_formula_headers or [])
    for col, header in enumerate(all_headers, start=1):
        ws.cell(row=3, column=col).value = header
    for row in range(4, 4 + data_row_count):
        for col in range(1, len(all_headers) + 1):
            ws.cell(row=row, column=col).number_format = "#,##0"
        if extra_formula_headers:
            formula_col = len(headers) + 1
            # Row-relative structured-reference formula (as the real APP_2
            # template's 広告売上 column uses "[#This Row]"/"@" syntax) --
            # unlike a plain "A5"-style reference, this is safe to copy
            # verbatim into a newly inserted row.
            ws.cell(row=row, column=formula_col).value = '=IFERROR([@広告表示数]*2,"-")'
    last_col_letter = ws.cell(row=3, column=len(all_headers)).column_letter
    table = Table(displayName=f"table_{sheet_name}", ref=f"A3:{last_col_letter}{3 + data_row_count}")
    table.tableColumns = [TableColumn(id=i, name=h) for i, h in enumerate(all_headers, start=1)]
    ws.add_table(table)
    return ws


class DateHelperTests(_RestoringTestCase):
    def test_previous_month_first(self):
        self.assertEqual(report.previous_month_first(date(2026, 9, 15)), date(2026, 8, 1))

    def test_parse_target_month_defaults_to_previous_month(self):
        self.assertEqual(report.parse_target_month(None, today=date(2026, 9, 1)), date(2026, 8, 1))

    def test_parse_target_month_rejects_non_first_day(self):
        with self.assertRaises(report.AdRevenueReportError) as ctx:
            report.parse_target_month("2026-08-15")
        self.assertEqual(ctx.exception.code, "invalid_target_month")

    def test_parse_target_month_rejects_bad_format(self):
        with self.assertRaises(report.AdRevenueReportError):
            report.parse_target_month("2026/08/01")


class OutputFileNameTests(_RestoringTestCase):
    def test_video_reward_file_name(self):
        name = report.output_file_name("video-reward", date(2026, 8, 1), date(2026, 9, 1))
        self.assertEqual(name, "J+_動画リワード広告売上_2026年08月期_260901.xlsx")

    def test_app2_file_name(self):
        name = report.output_file_name("app2", date(2026, 8, 1), date(2026, 9, 1))
        self.assertEqual(name, "J+_広告売上_APP_2_2026年08月期_260901.xlsx")

    def test_web_file_name(self):
        name = report.output_file_name("web", date(2026, 8, 1), date(2026, 9, 1))
        self.assertEqual(name, "J+_広告売上_WEB_2026年08月期_260901.xlsx")


class DefaultTemplateFileIdTests(_RestoringTestCase):
    """The official templates システム管理室 placed in the "広告売上" Drive
    folder (2026-09-18) are hardcoded as defaults so `generate` works with no
    env configuration at all -- matching thermae/plus's convention -- while
    each report type's env var can still override it independently."""

    def test_defaults_resolve_to_official_file_ids_with_no_env_config(self):
        with mock.patch.dict("os.environ", {}, clear=True):
            self.assertEqual(report.template_file_id("video-reward"), "1MtCimmJ9MjEjd0XW977kKME1sr82OP38")
            self.assertEqual(report.template_file_id("app2"), "1Q1d4Mp-5iQFBwo3wnzmUa1mQPxPdvL1L")
            self.assertEqual(report.template_file_id("web"), "1LJ72durOS1ekPO79GtDFdrJX23Vr8je-")

    def test_env_var_overrides_default_independently_per_report_type(self):
        with mock.patch.dict(
            "os.environ", {"AD_REVENUE_WEB_TEMPLATE_FILE_ID": "override-web-id"}, clear=True
        ):
            self.assertEqual(report.template_file_id("web"), "override-web-id")
            # other report types are unaffected by web's override
            self.assertEqual(report.template_file_id("video-reward"), "1MtCimmJ9MjEjd0XW977kKME1sr82OP38")
            self.assertEqual(report.template_file_id("app2"), "1Q1d4Mp-5iQFBwo3wnzmUa1mQPxPdvL1L")


class SheetCellParsingTests(_RestoringTestCase):
    def test_parse_iso_string_month(self):
        self.assertEqual(report._parse_sheet_month_cell("2026-08-01"), date(2026, 8, 1))

    def test_parse_serial_number_month(self):
        # Sheets serial date for 2026-08-01 (epoch 1899-12-30)
        serial = (date(2026, 8, 1) - date(1899, 12, 30)).days
        self.assertEqual(report._parse_sheet_month_cell(serial), date(2026, 8, 1))

    def test_parse_unparseable_month_fails_closed(self):
        with self.assertRaises(report.AdRevenueReportError) as ctx:
            report._parse_sheet_month_cell("not-a-date")
        self.assertEqual(ctx.exception.code, "unparseable_sheet_month")

    def test_parse_amount_with_commas(self):
        self.assertEqual(report._parse_sheet_amount_cell("14,017,945"), 14017945)

    def test_parse_amount_plain_number(self):
        self.assertEqual(report._parse_sheet_amount_cell(9789547), 9789547)

    def test_parse_amount_blank_is_none(self):
        self.assertIsNone(report._parse_sheet_amount_cell(""))
        self.assertIsNone(report._parse_sheet_amount_cell(None))
        self.assertIsNone(report._parse_sheet_amount_cell("-"))


class WriteSummaryTotalTests(_RestoringTestCase):
    def test_writes_first_total_when_only_one_block(self):
        wb = Workbook()
        wb.remove(wb.active)
        ws = _build_summary_sheet(wb)
        report.write_summary_total(ws, 9789547)
        self.assertEqual(ws["C4"].value, 9789547)

    def test_writes_first_total_when_multiple_total_blocks_exist(self):
        """Regression for 動画リワード's サマリ sheet, which has three separate
        "総計" rows (revenue / coin-consumption / author payout) -- only the
        first (topmost) one is the revenue cell this module should touch."""
        wb = Workbook()
        wb.remove(wb.active)
        ws = _build_summary_sheet(wb, with_multiple_total_blocks=True)
        report.write_summary_total(ws, 14017945)
        self.assertEqual(ws["C4"].value, 14017945)
        # the other two "総計" rows (coin consumption, author payout) are untouched
        self.assertEqual(ws["C9"].value, "#REF!")
        self.assertEqual(ws["C12"].value, "-")

    def test_missing_total_label_fails_closed(self):
        wb = Workbook()
        wb.remove(wb.active)
        ws = wb.create_sheet(report.SUMMARY_SHEET)
        ws["A1"] = "no total here"
        with self.assertRaises(report.AdRevenueReportError) as ctx:
            report.write_summary_total(ws, 100)
        self.assertEqual(ctx.exception.code, "summary_total_cell_not_found")


class WriteDetailSheetTests(_RestoringTestCase):
    def test_writes_rows_into_existing_capacity(self):
        wb = Workbook()
        wb.remove(wb.active)
        ws = _build_detail_sheet(wb, "全体", report.AD_VIEW_DETAIL_HEADERS, data_row_count=3)
        rows = [
            {
                "コンテンツID_Raise": f"ec{i}",
                "コンテンツID": i,
                "コンテンツ名": f"content-{i}",
                "JDCN": f"jdcn-{i}",
                "広告表示数": i * 100,
                "作品名": "work",
                "コミックスJDCN": "comic-jdcn",
                "コミックス巻数": 1,
                "タイトルID": 999,
                "デジタルタイトル名": "digital-title",
            }
            for i in range(1, 4)
        ]
        report.write_detail_sheet(ws, report.AD_VIEW_DETAIL_HEADERS, rows)

        self.assertEqual(ws.cell(row=4, column=1).value, "ec1")
        self.assertEqual(ws.cell(row=4, column=5).value, 100)
        self.assertEqual(ws.cell(row=6, column=1).value, "ec3")

    def test_grows_table_when_more_rows_than_template(self):
        wb = Workbook()
        wb.remove(wb.active)
        ws = _build_detail_sheet(wb, "全体", report.AD_VIEW_DETAIL_HEADERS, data_row_count=2)
        rows = [
            {h: (f"v{i}-{h}" if h != "広告表示数" else i * 10) for h in report.AD_VIEW_DETAIL_HEADERS}
            for i in range(1, 6)
        ]
        report.write_detail_sheet(ws, report.AD_VIEW_DETAIL_HEADERS, rows)

        table = ws.tables["table_全体"]
        self.assertEqual(table.ref, "A3:J8")
        self.assertEqual(ws.cell(row=8, column=5).value, 50)

    def test_shrinks_table_when_fewer_rows_than_template(self):
        wb = Workbook()
        wb.remove(wb.active)
        ws = _build_detail_sheet(wb, "全体", report.AD_VIEW_DETAIL_HEADERS, data_row_count=5)
        rows = [
            {h: (f"v{i}-{h}" if h != "広告表示数" else i) for h in report.AD_VIEW_DETAIL_HEADERS}
            for i in range(1, 3)
        ]
        report.write_detail_sheet(ws, report.AD_VIEW_DETAIL_HEADERS, rows)

        table = ws.tables["table_全体"]
        self.assertEqual(table.ref, "A3:J5")

    def test_formula_column_not_overwritten_and_survives_row_growth(self):
        """APP_2's 全体 sheet has a 広告売上 formula column this module never
        writes to; growing the table must copy that formula (not just style)
        into new rows so it keeps calculating."""
        wb = Workbook()
        wb.remove(wb.active)
        ws = _build_detail_sheet(
            wb, "全体", report.AD_VIEW_DETAIL_HEADERS, data_row_count=2, extra_formula_headers=["広告売上"]
        )
        rows = [
            {h: (f"v{i}-{h}" if h != "広告表示数" else i) for h in report.AD_VIEW_DETAIL_HEADERS} for i in range(1, 4)
        ]
        report.write_detail_sheet(ws, report.AD_VIEW_DETAIL_HEADERS, rows)

        formula_col = len(report.AD_VIEW_DETAIL_HEADERS) + 1
        # new row 6 (3rd data row) must have inherited the row-relative formula text
        self.assertEqual(ws.cell(row=6, column=formula_col).value, '=IFERROR([@広告表示数]*2,"-")')

    def test_missing_required_header_fails_closed(self):
        wb = Workbook()
        wb.remove(wb.active)
        ws = _build_detail_sheet(wb, "全体", ("コンテンツID_Raise", "コンテンツ名"), data_row_count=1)
        with self.assertRaises(report.AdRevenueReportError) as ctx:
            report.write_detail_sheet(ws, report.AD_VIEW_DETAIL_HEADERS, [{}])
        self.assertEqual(ctx.exception.code, "missing_detail_headers")


class CreateAdRevenueWorkbookTests(_RestoringTestCase):
    def _build_workbook(self, report_type):
        wb = Workbook()
        wb.remove(wb.active)
        _build_summary_sheet(wb)
        spec = report.REPORT_SPECS[report_type]
        for sheet_name in spec.detail_sheets:
            headers = spec.detail_headers
            if sheet_name == spec.zentai_sheet and spec.zentai_extra_headers:
                headers = headers + spec.zentai_extra_headers
            _build_detail_sheet(wb, sheet_name, headers, data_row_count=2)
        for work_spec in spec.work_summary_sheets:
            _build_detail_sheet(wb, work_spec.sheet_name, work_spec.headers, data_row_count=1)
        return wb

    def test_video_reward_end_to_end(self):
        """video-reward writes iOS, Android AND 全体 -- confirmed against the
        official template (2026-09-18): its サマリ sheet's コイン消費数 block
        sums live from 話データ_iOS[コイン消費数]/話データ_Android[コイン消費数],
        so leaving those two sheets blank would silently zero out that
        breakdown even though the total revenue cell (written directly,
        independent of those formulas) would still look correct."""
        wb = self._build_workbook("video-reward")
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)

        def _row(coin_count):
            return {
                "コンテンツID_Raise": "ec54987",
                "コンテンツID": 54987,
                "コンテンツ名": "[Ver2.01]title",
                "JDCN": "08X10000000038068100",
                "コイン消費数": coin_count,
                "作品名": "work",
                "コミックスJDCN": "comic-jdcn",
                "コミックス巻数": 1,
            }

        rows = {
            "iOS": [_row(455)],
            "Android": [_row(270)],
            "全体": [_row(725)],
        }
        with tempfile.TemporaryDirectory() as tmp_dir:
            output_path = Path(tmp_dir) / "out.xlsx"
            result = report.create_ad_revenue_workbook(
                report_type="video-reward",
                template_path=buf,
                output_path=output_path,
                revenue_yen=14017945,
                detail_rows=rows,
            )
            self.assertEqual(result["detail_row_count"], 3)

            from openpyxl import load_workbook

            reloaded = load_workbook(output_path)
        self.assertEqual(reloaded[report.SUMMARY_SHEET]["C4"].value, 14017945)
        self.assertEqual(reloaded["iOS"].cell(row=4, column=5).value, 455)
        self.assertEqual(reloaded["Android"].cell(row=4, column=5).value, 270)
        self.assertEqual(reloaded["全体"].cell(row=4, column=5).value, 725)

    def test_missing_summary_sheet_fails_closed(self):
        wb = Workbook()
        wb.remove(wb.active)
        spec = report.REPORT_SPECS["web"]
        for sheet_name in spec.detail_sheets:
            _build_detail_sheet(wb, sheet_name, spec.detail_headers, data_row_count=1)
        import io

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)

        with self.assertRaises(report.AdRevenueReportError) as ctx:
            report.create_ad_revenue_workbook(
                report_type="web",
                template_path=buf,
                output_path=io.BytesIO(),
                revenue_yen=1,
                detail_rows={},
            )
        self.assertEqual(ctx.exception.code, "summary_sheet_not_found")


class RunCoinContentQueryGoldenMasterTests(_RestoringTestCase):
    """Locks in the exact filter conditions derived from the 2026-08 Golden
    Master (docs/jumpplus-ad-revenue-report.md): without all four of
    service_name / ws_ex_is_return_reward_video_ad_coin / purchase_type /
    app_pf, plus the positive-sum exclusion, the totals and row counts don't
    match the real production file for that month."""

    def _run_and_capture_query(self, app_pfs):
        client = mock.Mock()
        query_job = mock.Mock()
        query_job.to_dataframe.return_value.to_dict.return_value = []
        client.query.return_value = query_job
        with mock.patch.object(report.bigquery, "Client", return_value=client):
            report.run_coin_content_query(project_id="p", target_month=date(2026, 8, 1), app_pfs=app_pfs)
        call = client.query.call_args
        return call.args[0], call.kwargs["job_config"]

    def test_query_text_includes_all_golden_master_conditions(self):
        query_text, _ = self._run_and_capture_query(["iOS", "And"])
        self.assertIn("ws_ex_is_return_reward_video_ad_coin = true", query_text)
        self.assertIn("purchase_type = 'episode'", query_text)
        self.assertIn("app_pf in unnest(@app_pfs)", query_text)
        self.assertIn("having reward_video_ad_coin_count > 0", query_text)

    def test_app_pfs_parameter_is_passed_through(self):
        _, job_config = self._run_and_capture_query(["iOS"])
        array_param = next(p for p in job_config.query_parameters if p.name == "app_pfs")
        self.assertEqual(array_param.values, ["iOS"])

    def test_query_text_fetches_ex_work_name_alongside_work_title(self):
        """作品別 sheet's Golden Master grouping (798 works) matches
        ex_work_name, not work_title (804 distinct values) -- confirmed
        against the real 2026-08 production file. work_title is still
        fetched for debugging only; _coin_content_row_to_detail must build
        the report's own 作品名 from ex_work_name."""
        query_text, _ = self._run_and_capture_query(["iOS", "And"])
        self.assertIn("any_value(ex_work_name) as ex_work_name", query_text)
        self.assertIn("any_value(work_title) as work_title", query_text)


class CoinContentRowToDetailTests(_RestoringTestCase):
    def test_work_name_column_comes_from_ex_work_name_not_work_title(self):
        record = {
            "prefixed_id": "ec1",
            "content_id": 1,
            "name": "content-1",
            "jdcn": "j1",
            "reward_video_ad_coin_count": 100,
            "work_title": "ONE PIECE　第1部",
            "ex_work_name": "ONE PIECE",
            "ex_comics_jdcn": None,
            "ex_episode_package_no": None,
        }
        detail = report._coin_content_row_to_detail(record)
        self.assertEqual(detail["作品名"], "ONE PIECE")

    def test_golden_master_ex_work_name_rollups_2026_08(self):
        """Locks in the real 2026-08 production file's per-work totals for
        works whose work_title variants collapse under ex_work_name (e.g.
        ONE PIECE's 第1部/第2部/第3部 all roll up to "ONE PIECE") --
        confirmed directly against the real Golden Master file, independent
        of the 798-vs-804 group-count investigation itself."""
        # (work_title, ex_work_name, coin_count) -- content-level rows that
        # collapse into one 作品別 group once grouped by ex_work_name.
        rows = [
            ("ONE PIECE　第1部", "ONE PIECE", 5_000_000),
            ("ONE PIECE　第2部", "ONE PIECE", 4_722_390),
            ("ONE PIECE　第3部", "ONE PIECE", 3_000_000),
            ("チェンソーマン 第一部", "チェンソーマン", 3_000_000),
            ("チェンソーマン 第二部", "チェンソーマン", 1_956_500),
            ("キン肉マン (38巻以降～、週プレ連載シリーズ)", "キン肉マン", 1_361_255),
            ("After World", "終末のハーレム", 208_920),
            ("奴隷遵戲　GUREN", "奴隷遵戲", 57_640),
            ("天神-TENJIN- イーグルネスト", "天神―TENJIN―", 23_690),
            ("放課後ましまし俱楽部", "声優ましまし俱楽部", 840),
        ]
        detail_rows = [
            report._coin_content_row_to_detail(
                {
                    "prefixed_id": f"ec{i}",
                    "content_id": i,
                    "name": f"content-{i}",
                    "jdcn": f"j{i}",
                    "reward_video_ad_coin_count": coin,
                    "work_title": work_title,
                    "ex_work_name": ex_work_name,
                    "ex_comics_jdcn": None,
                    "ex_episode_package_no": None,
                }
            )
            for i, (work_title, ex_work_name, coin) in enumerate(rows, start=1)
        ]

        summary = report.work_summaries.build_video_reward_work_summary(detail_rows, revenue_yen=14017945)
        by_name = {r["作品名"]: r["コイン消費数"] for r in summary}

        expected = {
            "ONE PIECE": 12_722_390,
            "チェンソーマン": 4_956_500,
            "キン肉マン": 1_361_255,
            "終末のハーレム": 208_920,
            "奴隷遵戲": 57_640,
            "天神―TENJIN―": 23_690,
            "声優ましまし俱楽部": 840,
        }
        for name, coin in expected.items():
            self.assertEqual(by_name[name], coin, f"{name} mismatch")
        # every work_title variant collapsed -- 10 content rows -> 7 works
        self.assertEqual(len(summary), 7)


class FetchDetailRowsDispatchTests(_RestoringTestCase):
    """Locks in each report type's mapping to BigQuery queries and detail
    sheets -- this is exactly the kind of thing that silently regresses (see
    video-reward's missing iOS/Android sheets, caught only by opening the
    official template directly, not by any prior test)."""

    def test_video_reward_queries_ios_android_and_combined(self):
        with mock.patch.object(report, "run_coin_content_query") as mock_query:
            mock_query.side_effect = [
                [{"content_id": 1, "reward_video_ad_coin_count": 455}],
                [{"content_id": 1, "reward_video_ad_coin_count": 270}],
                [{"content_id": 1, "reward_video_ad_coin_count": 725}],
            ]
            result = report.fetch_detail_rows(
                project_id="p", report_type="video-reward", target_month=date(2026, 8, 1)
            )

        self.assertEqual(set(result.keys()), {"iOS", "Android", "全体"})
        self.assertEqual(result["iOS"][0]["コイン消費数"], 455)
        self.assertEqual(result["Android"][0]["コイン消費数"], 270)
        self.assertEqual(result["全体"][0]["コイン消費数"], 725)
        calls = mock_query.call_args_list
        self.assertEqual(calls[0].kwargs["app_pfs"], ["iOS"])
        self.assertEqual(calls[1].kwargs["app_pfs"], ["And"])
        # combined 全体 sheet: iOS+And only, matching the Golden Master total
        # (365,253,010) -- video-reward never includes Web (always 0 anyway).
        self.assertEqual(calls[2].kwargs["app_pfs"], ["iOS", "And"])

    def test_app2_queries_ios_android_and_combined(self):
        with mock.patch.object(report, "run_ad_view_query") as mock_query:
            mock_query.side_effect = [
                [{"id": 1, "ad_view_count": 36680}],
                [{"id": 1, "ad_view_count": 10175}],
                [{"id": 1, "ad_view_count": 46855}],
            ]
            result = report.fetch_detail_rows(project_id="p", report_type="app2", target_month=date(2026, 8, 1))

        self.assertEqual(set(result.keys()), {"iOS", "Android", "全体"})
        calls = mock_query.call_args_list
        self.assertEqual(calls[0].kwargs["app_pfs"], ["iOS"])
        self.assertEqual(calls[1].kwargs["app_pfs"], ["And"])
        self.assertEqual(calls[2].kwargs["app_pfs"], ["iOS", "And"])

    def test_web_queries_only_web_platform_into_zentai_sheet(self):
        with mock.patch.object(report, "run_ad_view_query") as mock_query:
            mock_query.return_value = [{"id": 1, "ad_view_count": 21668}]
            result = report.fetch_detail_rows(project_id="p", report_type="web", target_month=date(2026, 8, 1))

        self.assertEqual(set(result.keys()), {"全体"})
        mock_query.assert_called_once()
        self.assertEqual(mock_query.call_args.kwargs["app_pfs"], ["Web"])

    def test_report_types_write_disjoint_firestore_and_revenue_columns(self):
        """Cross-report independence: each report_type has its own revenue
        column and its own Firestore claim key prefix (enforced by
        _ad_revenue_scheduled_run_id in app.py using f"{report_type}-{month}"),
        so one report waiting never blocks another from being generated."""
        columns = {spec.revenue_column for spec in report.REPORT_SPECS.values()}
        self.assertEqual(len(columns), len(report.REPORT_SPECS))


class AddWorkSummaryRowsTests(_RestoringTestCase):
    """Orchestration around ad_revenue_work_summaries.py: populates 全体's
    own 広告売上/広告売上_原資50/作品ID and the 作品別/作品別_2 replacements
    for what used to be Power Query."""

    def _zentai_row(self, i, *, work_id=None):
        return {
            "コンテンツID_Raise": f"ec{i}",
            "コンテンツID": i,
            "コンテンツ名": f"content-{i}",
            "JDCN": f"jdcn-{i}",
            "広告表示数": i * 10,
            "作品名": f"work-{i}",
            "コミックスJDCN": "comic-jdcn",
            "コミックス巻数": 1,
            "タイトルID": 900 + i,
            "デジタルタイトル名": f"digital-title-{i}",
            "作品ID": work_id if work_id is not None else i,
        }

    def test_app2_populates_zentai_and_both_work_summary_sheets(self):
        zentai = [self._zentai_row(1), self._zentai_row(2)]
        detail_rows = {"iOS": [], "Android": [], "全体": zentai}

        result = report.add_work_summary_rows(report_type="app2", revenue_yen=1000, detail_rows=detail_rows)

        self.assertIn("広告売上", result["全体"][0])
        self.assertIn("広告売上_原資50", result["全体"][0])
        self.assertEqual(set(result.keys()), {"iOS", "Android", "全体", "作品別", "作品別_2"})
        self.assertEqual(len(result["作品別"]), 2)
        self.assertEqual(len(result["作品別_2"]), 2)
        # revenue is split proportionally to 広告表示数 (10 and 20) -- total
        # ad revenue across 作品別 groups must sum back to the input revenue.
        total = sum((r["広告売上"] for r in result["作品別"]), start=type(result["作品別"][0]["広告売上"])(0))
        self.assertEqual(round(total), 1000)

    def test_web_populates_zentai_and_single_work_summary_sheet(self):
        zentai = [self._zentai_row(1), self._zentai_row(2)]
        detail_rows = {"全体": zentai}

        result = report.add_work_summary_rows(report_type="web", revenue_yen=500, detail_rows=detail_rows)

        self.assertIn("広告売上", result["全体"][0])
        self.assertNotIn("広告売上_原資50", result["全体"][0])
        self.assertEqual(set(result.keys()), {"全体", "作品別"})
        self.assertEqual(len(result["作品別"]), 2)

    def test_video_reward_adds_only_work_summary_sheet(self):
        zentai = [
            {"作品名": "A", "コイン消費数": 300},
            {"作品名": "B", "コイン消費数": 700},
        ]
        detail_rows = {"iOS": [], "Android": [], "全体": zentai}

        result = report.add_work_summary_rows(
            report_type="video-reward", revenue_yen=14017945, detail_rows=detail_rows
        )

        self.assertEqual(set(result.keys()), {"iOS", "Android", "全体", "作品別"})
        # 全体 rows themselves are untouched (no new keys added, unlike app2/web)
        self.assertEqual(set(result["全体"][0].keys()), {"作品名", "コイン消費数"})
        self.assertEqual(len(result["作品別"]), 2)

    def test_does_not_mutate_input_dict(self):
        zentai = [self._zentai_row(1)]
        detail_rows = {"iOS": [], "Android": [], "全体": zentai}
        original_zentai_row_keys = set(zentai[0].keys())

        report.add_work_summary_rows(report_type="app2", revenue_yen=100, detail_rows=detail_rows)

        self.assertEqual(set(detail_rows["全体"][0].keys()), original_zentai_row_keys)
        self.assertNotIn("作品別", detail_rows)


class ReadinessTests(_RestoringTestCase):
    def _mock_client(self, *, detail_rows, revenue_row):
        """check_readiness always issues the detail-ready query first, then
        (only if ready) the confirmed-revenue query -- so a plain call-order
        side_effect list is enough and doesn't need to parse SQL text."""
        client = mock.Mock()
        detail_query_result = mock.Mock()
        detail_query_result.result.return_value = detail_rows
        revenue_query_result = mock.Mock()
        revenue_query_result.result.return_value = [revenue_row] if revenue_row is not None else []
        client.query.side_effect = [detail_query_result, revenue_query_result]
        return client

    def test_waiting_detail_when_no_detail_rows(self):
        client = self._mock_client(detail_rows=[], revenue_row=None)
        with mock.patch.object(report.bigquery, "Client", return_value=client):
            result = report.check_readiness(project_id="p", report_type="web", target_month=date(2026, 8, 1))
        self.assertEqual(result.status, "waiting_detail")

    def test_waiting_revenue_when_detail_ready_but_no_revenue(self):
        client = self._mock_client(detail_rows=[{"f": 1}], revenue_row=None)
        with mock.patch.object(report.bigquery, "Client", return_value=client):
            result = report.check_readiness(project_id="p", report_type="web", target_month=date(2026, 8, 1))
        self.assertEqual(result.status, "waiting_revenue")

    def test_ready_when_both_detail_and_revenue_present(self):
        client = self._mock_client(detail_rows=[{"f": 1}], revenue_row={"yen": 734139})
        with mock.patch.object(report.bigquery, "Client", return_value=client):
            result = report.check_readiness(project_id="p", report_type="web", target_month=date(2026, 8, 1))
        self.assertEqual(result.status, "ready")
        self.assertEqual(result.revenue_yen, 734139)

    def test_readiness_is_independent_per_report_type(self):
        """The exact scenario from the intake requirements: for the same
        target month, video-reward and app2 can be ready while web is still
        waiting -- one report type's missing confirmed value must never
        affect another's readiness."""
        target_month = date(2026, 8, 1)
        scenarios = {
            "video-reward": ({"f": 1}, {"yen": 14017945}, "ready"),
            "app2": ({"f": 1}, {"yen": 9789547}, "ready"),
            "web": ({"f": 1}, None, "waiting_revenue"),
        }
        for report_type, (detail_row, revenue_row, expected_status) in scenarios.items():
            client = self._mock_client(detail_rows=[detail_row], revenue_row=revenue_row)
            with mock.patch.object(report.bigquery, "Client", return_value=client):
                result = report.check_readiness(project_id="p", report_type=report_type, target_month=target_month)
            self.assertEqual(result.status, expected_status, msg=report_type)


class GenerateReportGuardTests(_RestoringTestCase):
    def test_unknown_report_type_rejected(self):
        with self.assertRaises(report.AdRevenueReportError) as ctx:
            report.generate_ad_revenue_report(project_id="p", report_type="bogus")
        self.assertEqual(ctx.exception.code, "unknown_report_type")

    def test_not_ready_raises_409(self):
        with mock.patch.object(
            report, "check_readiness", return_value=report.ReadinessResult(status="waiting_revenue")
        ):
            with self.assertRaises(report.AdRevenueReportError) as ctx:
                report.generate_ad_revenue_report(
                    project_id="p", report_type="web", target_month_text="2026-08-01"
                )
        self.assertEqual(ctx.exception.code, "not_ready")
        self.assertEqual(ctx.exception.status_code, 409)

    def test_template_not_configured_raises(self):
        """Templates now have real hardcoded defaults (see REPORT_SPECS), so
        this only fires when an operator explicitly overrides the env var to
        empty -- e.g. during an incident, to force-disable one report type
        without touching the others."""
        with mock.patch.object(
            report, "check_readiness", return_value=report.ReadinessResult(status="ready", revenue_yen=1)
        ), mock.patch.object(report, "fetch_detail_rows", return_value={"全体": []}), mock.patch.dict(
            "os.environ", {"AD_REVENUE_WEB_TEMPLATE_FILE_ID": ""}, clear=True
        ):
            with self.assertRaises(report.AdRevenueReportError) as ctx:
                report.generate_ad_revenue_report(
                    project_id="p", report_type="web", target_month_text="2026-08-01"
                )
        self.assertEqual(ctx.exception.code, "template_not_configured")


class SheetSyncTests(_RestoringTestCase):
    def test_unexpected_header_fails_closed(self):
        with mock.patch("sheets_io.read_sheet_values", return_value=[["wrong", "header"]]):
            with self.assertRaises(report.AdRevenueReportError) as ctx:
                report.sync_ad_revenue_confirmed_values(project_id="p")
        self.assertEqual(ctx.exception.code, "unexpected_sheet_header")

    def test_empty_sheet_fails_closed(self):
        with mock.patch("sheets_io.read_sheet_values", return_value=[]):
            with self.assertRaises(report.AdRevenueReportError) as ctx:
                report.sync_ad_revenue_confirmed_values(project_id="p")
        self.assertEqual(ctx.exception.code, "sheet_empty")

    def test_merges_parsed_rows_into_bigquery(self):
        sheet_rows = [
            ["対象月", "動画リワード", "奥付広告", "WEB"],
            ["2026-08-01", 14017945, 9789547, 734139],
            ["2026-07-01", "15,872,954", "11,644,698", "1,414,528"],
            ["", "", "", ""],
        ]
        client = mock.Mock()
        with mock.patch("sheets_io.read_sheet_values", return_value=sheet_rows), mock.patch.object(
            report.bigquery, "Client", return_value=client
        ):
            result = report.sync_ad_revenue_confirmed_values(project_id="p")

        self.assertEqual(result, {"status": "ok", "row_count": 2})
        client.create_table.assert_called_once()
        client.query.assert_called_once()
        merge_call = client.query.call_args
        job_config = merge_call.kwargs["job_config"]
        array_param = job_config.query_parameters[0]
        self.assertEqual(len(array_param.values), 2)


if __name__ == "__main__":
    unittest.main()
