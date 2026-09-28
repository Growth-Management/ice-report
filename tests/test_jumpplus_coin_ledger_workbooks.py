import os
import tempfile
import unittest
import zipfile
from datetime import date
from decimal import Decimal
from pathlib import Path

from lxml import etree
from openpyxl import load_workbook

import jumpplus_coin_ledger_report as report
import jumpplus_coin_ledger_workbooks as workbooks
import xlsx_package_writer as pkg_writer
from tests import coin_ledger_fixtures as fx

M = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
TARGET = date(2026, 8, 1)

# Directory holding the real official templates (never committed): the env
# var, else the documented local working-copy location when it exists
# (docs/jumpplus-coin-ledger-report.md "テンプレート"). Machines without
# either (e.g. CI) skip only these tests.
DEFAULT_REAL_TEMPLATE_DIR = r"C:\temp\jumpplus-coin-ledger-templates"
REAL_TEMPLATE_DIR = os.environ.get("JUMPPLUS_COIN_LEDGER_TEMPLATE_DIR") or (
    DEFAULT_REAL_TEMPLATE_DIR if Path(DEFAULT_REAL_TEMPLATE_DIR).is_dir() else ""
)
REAL_TEMPLATE_NAMES = {
    "app": "【少年ジャンプ＋】消費コイン_yyyy年mm月期_yymmdd.xlsx",
    "web": "【少年ジャンプ＋】消費コイン_WEB_yyyy年mm月期_yymmdd.xlsx",
    "product": "話売商品_一覧_yymm.xlsx",
}


def _read_part(path, part):
    with zipfile.ZipFile(path) as zf:
        return zf.read(part)


def _cell_style(path, part, ref):
    root = etree.fromstring(_read_part(path, part))
    for c in root.iter(M + "c"):
        if c.get("r") == ref:
            return c.get("s")
    return None


def _formula_texts(path, part):
    root = etree.fromstring(_read_part(path, part))
    return {c.get("r"): c.find(M + "f").text for c in root.iter(M + "c") if c.find(M + "f") is not None}


class _Base(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self._tmp.cleanup)
        self.tmp = Path(self._tmp.name)
        self.source = fx.dummy_source()


class LedgerSheetTests(_Base):
    def _build(self, source=None):
        template = fx.build_app_template(self.tmp / "app_t.xlsx")
        out = self.tmp / "app.xlsx"
        workbooks.build_app_workbook(template_path=template, output_path=out, target_month=TARGET, source=source or self.source)
        return template, out

    def test_three_blocks_filled_from_ledger(self):
        _, out = self._build()
        ws = load_workbook(out)["出納"]
        by = {(r["app_pf"], r["coin_type"]): r for r in self.source.ledger}
        for offset, coin_type in enumerate(report.COIN_TYPES):
            ios, android = by[("iOS", coin_type)], by[("And", coin_type)]
            self.assertEqual(ws[f"B{4 + offset}"].value, int(ios["issue_coins_m"] + android["issue_coins_m"]))
            self.assertEqual(ws[f"D{4 + offset}"].value, int(ios["use_coins_m"] + android["use_coins_m"]))
            self.assertEqual(ws[f"C{16 + offset}"].value, int(ios["carried_over_coins_m"]))
            self.assertEqual(ws[f"E{16 + offset}"].value, int(ios["cancellation_coins_m"]))
            self.assertEqual(ws[f"D{28 + offset}"].value, int(android["use_coins_m"]))

    def test_formulas_and_styles_are_preserved(self):
        template, out = self._build()
        part = "xl/worksheets/sheet1.xml"
        self.assertEqual(_formula_texts(template, part), _formula_texts(out, part))
        for ref in ("B4", "E9", "C16", "D33", "F4", "B10"):
            self.assertEqual(_cell_style(template, part, ref), _cell_style(out, part, ref))

    def test_web_block_uses_web_platform_only(self):
        template = fx.build_web_template(self.tmp / "web_t.xlsx")
        out = self.tmp / "web.xlsx"
        workbooks.build_web_workbook(template_path=template, output_path=out, target_month=TARGET, source=self.source)
        ws = load_workbook(out)["出納"]
        web = {r["coin_type"]: r for r in self.source.ledger if r["app_pf"] == "Web"}
        self.assertEqual(ws["D4"].value, int(web["pay_coin"]["use_coins_m"]))
        self.assertEqual(ws["B9"].value, int(web["reward_video_ad_coin"]["issue_coins_m"]))

    def test_balance_must_match_template_formula(self):
        self.source.ledger[0]["balance_coins_m"] += 1
        with self.assertRaises(report.CoinLedgerReportError) as ctx:
            self._build()
        self.assertEqual(ctx.exception.code, "ledger_balance_mismatch")

    def test_missing_grain_fails_closed(self):
        self.source.ledger = [
            r for r in self.source.ledger if not (r["app_pf"] == "iOS" and r["coin_type"] == "pay_gift_coin")
        ]
        with self.assertRaises(report.CoinLedgerReportError) as ctx:
            self._build()
        self.assertEqual(ctx.exception.code, "ledger_grain_mismatch")

    def test_layout_change_in_template_fails_closed(self):
        template = fx.build_app_template(self.tmp / "app_t.xlsx")
        wb = load_workbook(template)
        wb["出納"]["A5"] = "別のコイン"
        wb.save(template)
        with self.assertRaises(report.CoinLedgerReportError) as ctx:
            workbooks.build_app_workbook(template_path=template, output_path=self.tmp / "o.xlsx", target_month=TARGET, source=self.source)
        self.assertEqual(ctx.exception.code, "ledger_template_layout_mismatch")


class SummaryTests(_Base):
    def test_summary_rows_group_by_work_and_sort_by_total(self):
        rows = [
            {"ex_work_name": "A", "purchase_type": "episode", "total_use_coins": Decimal(10), "pay_coins_total": Decimal(10)},
            {"ex_work_name": "B", "purchase_type": "book", "total_use_coins": Decimal(30), "free_ad_coins_total": Decimal(30), "ex_comic_type": "X"},
            {"ex_work_name": "A", "purchase_type": "book", "total_use_coins": Decimal(5), "pay_gift_coins_total": Decimal(5)},
        ]
        out = workbooks.build_summary_rows(rows, unit="コイン", works=["A", "B", "C"], with_type=True)
        self.assertEqual([r["作品名"] for r in out], ["B", "A", "C"])
        a = out[1]
        self.assertEqual(a["コイン消費合計"], 15)
        self.assertEqual(a["話\n有償コイン消費数"], 10)
        self.assertEqual(a["巻\n贈答コイン消費数"], 5)
        self.assertEqual(out[0]["巻\n広告コイン消費数"], 30)
        self.assertEqual(out[0]["種別"], "X")
        self.assertEqual(out[2]["コイン消費合計"], 0)

    def test_app_summary_sheets_share_work_list_and_keep_totals_row(self):
        template = fx.build_app_template(self.tmp / "app_t.xlsx")
        out = self.tmp / "app.xlsx"
        counts = workbooks.build_app_workbook(template_path=template, output_path=out, target_month=TARGET, source=self.source)
        self.assertEqual(counts["サマリ"], counts["サマリ (Apple)"])
        self.assertEqual(counts["サマリ"], counts["サマリ (Google)"])
        wb = load_workbook(out)
        for sheet_name in ("サマリ", "サマリ (Apple)", "サマリ (Google)"):
            ws = wb[sheet_name]
            table = list(ws.tables.values())[0]
            last = int(table.ref.split(":")[1][1:])
            self.assertEqual(last, 3 + counts[sheet_name] + 1, sheet_name)  # shrink from 30 blank rows
            self.assertEqual(ws.cell(last, 1).value, "総計")
            self.assertTrue(str(ws.cell(last, 2).value).startswith("=SUBTOTAL(109,"))
        total_rows = [r for r in self.source.content if int(r["app_id"]) in (31, 32)]
        written = sum(wb["サマリ"].cell(r, 2).value for r in range(4, 4 + counts["サマリ"]))
        self.assertEqual(written, int(sum(r["total_use_coins"] for r in total_rows)))


class DetailTests(_Base):
    def test_detail_row_contract(self):
        record = self.source.content[1]
        episode = workbooks.detail_row(record, unit="コイン", purchase_type="episode")
        self.assertEqual(tuple(episode), workbooks.detail_headers("コイン", "episode"))
        self.assertEqual(episode["コンテンツID"], record["v2_content_id_token"])
        self.assertEqual(episode["備考"], "-")
        self.assertEqual(episode["配信開始日"], record["ex_sales_start_date"].isoformat())
        self.assertIsInstance(episode["消費コイン"], int)
        book = workbooks.detail_row(record, unit="ポイント", purchase_type="book")
        self.assertEqual((book["雑誌"], book["種別"]), ("-", "-"))
        self.assertIn("価格（ポイント）", book)

    def test_detail_tables_grow_and_split_by_platform_and_type(self):
        template = fx.build_app_template(self.tmp / "app_t.xlsx")
        out = self.tmp / "app.xlsx"
        counts = workbooks.build_app_workbook(template_path=template, output_path=out, target_month=TARGET, source=self.source)
        for sheet_name, app_id, purchase_type in workbooks.APP_DETAIL_SHEETS:
            expected = len(report.content_rows_for(self.source.content, app_id=app_id, purchase_type=purchase_type))
            self.assertEqual(counts[sheet_name], expected)
        wb = load_workbook(out)
        ws = wb["有料話消費コイン（Apple）"]
        self.assertEqual(list(ws.tables.values())[0].ref, f"A3:R{3 + counts['有料話消費コイン（Apple）']}")
        self.assertEqual(ws.freeze_panes, "E4")
        self.assertEqual(ws.sheet_view.zoomScale, 80)

    def test_empty_sheet_keeps_one_blank_row(self):
        source = fx.dummy_source()
        source.content = [r for r in source.content if not (int(r["app_id"]) == 101 and r["purchase_type"] == "book")]
        template = fx.build_web_template(self.tmp / "web_t.xlsx")
        out = self.tmp / "web.xlsx"
        workbooks.build_web_workbook(template_path=template, output_path=out, target_month=TARGET, source=source)
        ws = load_workbook(out)["有料巻消費ポイント"]
        self.assertEqual(list(ws.tables.values())[0].ref, "A3:R4")
        self.assertIsNone(ws["A4"].value)

    def test_missing_template_header_fails_closed(self):
        template = fx.build_app_template(self.tmp / "app_t.xlsx")
        wb = load_workbook(template)
        wb["有料巻消費コイン（Google）"]["Q3"] = "別の列"
        wb.save(template)
        with self.assertRaises(report.CoinLedgerReportError) as ctx:
            workbooks.build_app_workbook(template_path=template, output_path=self.tmp / "o.xlsx", target_month=TARGET, source=self.source)
        self.assertEqual(ctx.exception.code, "template_header_missing")


class ProductTests(_Base):
    def test_product_master_mapping(self):
        template = fx.build_product_template(self.tmp / "p_t.xlsx")
        out = self.tmp / "p.xlsx"
        counts = workbooks.build_product_workbook(template_path=template, output_path=out, target_month=TARGET, source=self.source)
        self.assertEqual(counts, {"file": len(self.source.product_master)})
        ws = load_workbook(out)["file"]
        self.assertEqual(list(ws.tables.values())[0].ref, f"A1:J{1 + len(self.source.product_master)}")
        self.assertEqual(ws.freeze_panes, "A2")
        first = self.source.product_master[0]
        header = [c.value for c in ws[1]]
        row = dict(zip(header, [c.value for c in ws[2]]))
        self.assertEqual(row["コンテンツID_Raise"], first["prefixed_id"])
        self.assertEqual(row["コンテンツID"], first["v2_content_id_token"])
        self.assertEqual(row["作品名"], first["work_title"])
        self.assertEqual(row["著者名"], first["author_name"])
        self.assertEqual(row["価格（コイン）"], first["price_in_coin"])
        self.assertEqual(row["配信開始日"], first["ex_sales_start_date"].isoformat())


class PackagePreservationTests(_Base):
    def test_untouched_parts_are_byte_identical_and_output_is_valid(self):
        template = fx.build_app_template(self.tmp / "app_t.xlsx")
        out = self.tmp / "app.xlsx"
        workbooks.build_app_workbook(template_path=template, output_path=out, target_month=TARGET, source=self.source)
        preserved = [n for n in pkg_writer.list_parts(template) if n.startswith(("xl/styles", "xl/theme", "docProps"))]
        self.assertTrue(all(pkg_writer.validate_preserved_parts(template, out, preserved).values()))
        info = report.validate_xlsx_output(out)
        self.assertEqual(info["sheet_count"], 8)
        with zipfile.ZipFile(out) as zf:
            self.assertIsNone(zf.testzip())
            o_root = etree.fromstring(zf.read("xl/worksheets/sheet5.xml"))
        t_root = etree.fromstring(_read_part(template, "xl/worksheets/sheet5.xml"))
        for tag in ("sheetViews", "cols"):
            te, oe = t_root.find(M + tag), o_root.find(M + tag)
            self.assertEqual(etree.tostring(te) if te is not None else None, etree.tostring(oe) if oe is not None else None)
        calc = etree.fromstring(_read_part(out, "xl/workbook.xml")).find(M + "calcPr")
        self.assertEqual(calc.get("fullCalcOnLoad"), "1")


class PackageWriterHelperTests(_Base):
    def test_write_fixed_cells_refuses_formula_and_keeps_style(self):
        template = fx.build_app_template(self.tmp / "app_t.xlsx")
        pkg = pkg_writer.XlsxPackage.load(template)
        with self.assertRaises(pkg_writer.XlsxPackageError) as ctx:
            pkg_writer.write_fixed_cells(pkg, "出納", {"F4": 1})
        self.assertEqual(ctx.exception.code, "formula_cell_overwrite")
        pkg_writer.write_fixed_cells(pkg, "出納", {"B4": 123, "Z40": "new"})
        self.assertEqual(pkg_writer.read_cell_texts(pkg, "出納", ["B4", "Z40", "A4"]), {"B4": "123", "Z40": "new", "A4": "有償コイン"})

    def test_sync_filter_database_names(self):
        template = fx.build_web_template(self.tmp / "web_t.xlsx")
        pkg = pkg_writer.XlsxPackage.load(template)
        wb_root = pkg.xml("xl/workbook.xml")
        names = wb_root.find(M + "definedNames")
        if names is None:
            names = etree.SubElement(wb_root, M + "definedNames")
        for local_id, text in (("1", "サマリ!$A$3:$O$5"), ("2", "有料話消費ポイント!$A$3:$R$4")):
            el = etree.SubElement(names, M + "definedName", name="_xlnm._FilterDatabase", localSheetId=local_id, hidden="1")
            el.text = text
        pkg.set_xml("xl/workbook.xml", wb_root)
        pkg_writer.replace_detail_rows(pkg, "有料話消費ポイント", ("コンテンツID_Raise",), [{"コンテンツID_Raise": "x"}] * 5)
        self.assertEqual(sorted(pkg_writer.sync_filter_database_names(pkg)), ["サマリ", "有料話消費ポイント"])
        texts = [d.text for d in pkg.xml("xl/workbook.xml").iter(M + "definedName")]
        self.assertIn("有料話消費ポイント!$A$3:$R$8", texts)


@unittest.skipUnless(REAL_TEMPLATE_DIR, "official templates not found (set JUMPPLUS_COIN_LEDGER_TEMPLATE_DIR)")
class RealTemplateTests(_Base):
    def _template(self, key):
        path = Path(REAL_TEMPLATE_DIR) / REAL_TEMPLATE_NAMES[key]
        if not path.exists():
            self.skipTest(f"{path.name} not present")
        return path

    def _check(self, key, builder, sheet_count):
        template = self._template(key)
        out = self.tmp / f"{key}.xlsx"
        source = fx.dummy_source(contents_per_app=300, works=40)
        builder(template_path=template, output_path=out, target_month=TARGET, source=source)
        self.assertEqual(report.validate_xlsx_output(out)["sheet_count"], sheet_count)
        preserved = [n for n in pkg_writer.list_parts(template) if n.startswith(("xl/styles", "xl/theme", "xl/printerSettings", "docProps"))]
        self.assertTrue(all(pkg_writer.validate_preserved_parts(template, out, preserved).values()))
        with zipfile.ZipFile(template) as tz, zipfile.ZipFile(out) as oz:
            for name in (n for n in tz.namelist() if n.startswith("xl/worksheets/sheet")):
                t_root, o_root = etree.fromstring(tz.read(name)), etree.fromstring(oz.read(name))
                for tag in ("sheetViews", "cols", "mergeCells"):
                    te, oe = t_root.find(M + tag), o_root.find(M + tag)
                    self.assertEqual(etree.tostring(te) if te is not None else None, etree.tostring(oe) if oe is not None else None, (name, tag))
            for name in (n for n in tz.namelist() if n.startswith("xl/tables/")):
                t_root, o_root = etree.fromstring(tz.read(name)), etree.fromstring(oz.read(name))
                self.assertEqual(etree.tostring(t_root.find(M + "tableColumns")), etree.tostring(o_root.find(M + "tableColumns")))
        return out

    def test_real_app_template(self):
        out = self._check("app", workbooks.build_app_workbook, 8)
        wb = load_workbook(out, read_only=True)
        self.assertEqual(wb.sheetnames[0], "出納")
        wb.close()

    def test_real_web_template(self):
        self._check("web", workbooks.build_web_workbook, 4)

    def test_real_product_template(self):
        wb = load_workbook(self._template("product"), read_only=True)
        sheet_count = len(wb.sheetnames)
        wb.close()
        self._check("product", workbooks.build_product_workbook, sheet_count)


if __name__ == "__main__":
    unittest.main()
