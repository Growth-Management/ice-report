import io
import re
import sys
import unittest
import zipfile

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.filters import AutoFilter
from openpyxl.worksheet.table import Table, TableColumn

import plus_browser_point_sales_report as report

# Some sibling test modules (e.g. test_admin_report_definitions.py) replace
# sys.modules["openpyxl"] with an incomplete stub and never restore it, which
# breaks openpyxl's real Workbook.save() (it does `from openpyxl import
# __version__` lazily at call time) for any test that runs afterward in the
# same process via `python -m unittest discover`. Capture the real modules at
# import time here and reinstate them before saving, so this file's result
# does not depend on test execution order.
_REAL_OPENPYXL_MODULES = {
    name: mod for name, mod in sys.modules.items() if name == "openpyxl" or name.startswith("openpyxl.")
}


def _restore_real_openpyxl() -> None:
    sys.modules.update(_REAL_OPENPYXL_MODULES)

PRODUCT_LABELS = [report.PRICE_TO_POINT_LABEL[p] for p in sorted(report.PRICE_TO_POINT_LABEL)]
BASE_PAYMENT_CLASSES = [
    "1_credit(JCB/AMEX)",
    "2_credit(SBPS)",
    "3_credit(DINERS)",
    "4_credit(その他)",
    "5_docomo",
    "6_au",
    "7_softbank",
    "8_paypay",
    "9_rakuten",
    "10_suica",
]


def _build_template_workbook() -> Workbook:
    """A minimal but structurally faithful stand-in for the real Drive template.

    Mirrors the production template's two Excel Tables (10 data rows + a
    totals row that already contains the pre-existing SUBTOTAL structured
    references) without depending on the real, non-versioned template file.
    """
    wb = Workbook()
    wb.remove(wb.active)

    ws1 = wb.create_sheet(report.BY_PAYMENT_SHEET)
    ws1["A1"] = "決済分類"
    ws1["B1"] = "売上"
    for row in range(2, 12):
        ws1.cell(row=row, column=2).number_format = "#,##0"
    ws1["A12"] = "合計"
    ws1["B12"] = f"=SUBTOTAL(109,{report.BY_PAYMENT_TABLE}[売上])"
    ws1["B12"].number_format = "#,##0"
    table1 = Table(displayName=report.BY_PAYMENT_TABLE, ref="A1:B12")
    table1.tableColumns = [
        TableColumn(id=1, name="決済分類"),
        TableColumn(id=2, name="売上", totalsRowFunction="sum"),
    ]
    table1.totalsRowCount = 1
    table1.autoFilter = AutoFilter(ref="A1:B11")
    ws1.add_table(table1)

    ws2 = wb.create_sheet(report.BY_PAYMENT_PRODUCT_SHEET)
    headers = ["決済分類", *PRODUCT_LABELS, "TOTAL"]
    for col, header in enumerate(headers, start=1):
        ws2.cell(row=1, column=col).value = header
    for row in range(2, 12):
        for col in range(2, len(headers) + 1):
            ws2.cell(row=row, column=col).number_format = "#,##0"
    ws2.cell(row=12, column=1).value = "TOTAL"
    for col, label in enumerate(headers[1:], start=2):
        cell = ws2.cell(row=12, column=col)
        cell.value = f"=SUBTOTAL(109,{report.BY_PAYMENT_PRODUCT_TABLE}[{label}])"
        cell.number_format = "#,##0"
    table2 = Table(displayName=report.BY_PAYMENT_PRODUCT_TABLE, ref=f"A1:{get_column_letter(len(headers))}12")
    table2.tableColumns = [
        TableColumn(id=i, name=h, totalsRowFunction=("sum" if i > 1 else None))
        for i, h in enumerate(headers, start=1)
    ]
    table2.totalsRowCount = 1
    table2.autoFilter = AutoFilter(ref=f"A1:{get_column_letter(len(headers))}11")
    ws2.add_table(table2)

    return wb


def _make_data(payment_classes: list[str]) -> dict:
    prices = sorted(report.PRICE_TO_POINT_LABEL.keys())
    by_payment_total = {pc: (i + 1) * 1000 for i, pc in enumerate(payment_classes)}
    by_payment_product = {
        pc: {price: (i + 1) * 100 for price in prices} for i, pc in enumerate(payment_classes)
    }
    return {
        "payment_classes": payment_classes,
        "prices": prices,
        "by_payment_total": by_payment_total,
        "by_payment_product": by_payment_product,
        "grand_total": sum(by_payment_total.values()),
    }


def _save_and_reload(wb: Workbook):
    _restore_real_openpyxl()
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf, load_workbook(buf)


class WriteByPaymentProductSheetTests(unittest.TestCase):
    def test_total_formula_is_plain_a1_range_sum(self):
        """Test 1: the fix -- data-row TOTAL cells use =SUM(Bn:Kn)."""
        wb = _build_template_workbook()
        data = _make_data(BASE_PAYMENT_CLASSES)
        report.write_by_payment_product_sheet(wb[report.BY_PAYMENT_PRODUCT_SHEET], data)

        ws = wb[report.BY_PAYMENT_PRODUCT_SHEET]
        last_col_letter = get_column_letter(1 + len(data["prices"]))
        for offset in range(len(BASE_PAYMENT_CLASSES)):
            row = 2 + offset
            self.assertEqual(
                ws.cell(row=row, column=1 + len(data["prices"]) + 1).value,
                f"=SUM(B{row}:{last_col_letter}{row})",
            )

    def test_no_structured_reference_survives_in_saved_xml(self):
        """Test 2: the saved xlsx must not contain the table structured-reference
        syntax ("[@[") anywhere in a worksheet part -- that syntax is exactly
        what triggered Excel's repair-and-drop-formula behavior."""
        wb = _build_template_workbook()
        data = _make_data(BASE_PAYMENT_CLASSES)
        report.write_by_payment_sheet(wb[report.BY_PAYMENT_SHEET], data)
        report.write_by_payment_product_sheet(wb[report.BY_PAYMENT_PRODUCT_SHEET], data)

        buf, _ = _save_and_reload(wb)
        with zipfile.ZipFile(buf) as zf:
            worksheet_parts = [n for n in zf.namelist() if re.match(r"xl/worksheets/sheet\d+\.xml$", n)]
            self.assertTrue(worksheet_parts, "expected at least one worksheet XML part")
            for name in worksheet_parts:
                xml_text = zf.read(name).decode("utf-8")
                self.assertNotIn(
                    f"{report.BY_PAYMENT_PRODUCT_TABLE}[@",
                    xml_text,
                    f"found a structured-reference formula in {name}",
                )
                self.assertNotIn("[@[", xml_text, f"found table row-scoped reference syntax in {name}")

    def test_row_extension_keeps_total_formula_a1_ranges_in_sync(self):
        """Test 3: extending the table (11th/12th payment_class) keeps each new
        row's TOTAL formula pointing at that same row, not a stale one."""
        wb = _build_template_workbook()
        extended_classes = BASE_PAYMENT_CLASSES + ["11_new_method", "12_second_new_method"]
        data = _make_data(extended_classes)
        report.write_by_payment_product_sheet(wb[report.BY_PAYMENT_PRODUCT_SHEET], data)

        ws = wb[report.BY_PAYMENT_PRODUCT_SHEET]
        last_col_letter = get_column_letter(1 + len(data["prices"]))
        self.assertEqual(ws.cell(row=12, column=12).value, f"=SUM(B12:{last_col_letter}12)")
        self.assertEqual(ws.cell(row=13, column=12).value, f"=SUM(B13:{last_col_letter}13)")

    def test_totals_row_keeps_structured_subtotal_reference(self):
        """Test 4: the table's own totals row is untouched -- it keeps the
        pre-existing structured SUBTOTAL reference, which does not trigger the
        Excel repair behavior (verified on real Windows Excel)."""
        wb = _build_template_workbook()
        data = _make_data(BASE_PAYMENT_CLASSES)
        report.write_by_payment_product_sheet(wb[report.BY_PAYMENT_PRODUCT_SHEET], data)

        ws = wb[report.BY_PAYMENT_PRODUCT_SHEET]
        self.assertEqual(
            ws.cell(row=12, column=12).value,
            f"=SUBTOTAL(109,{report.BY_PAYMENT_PRODUCT_TABLE}[TOTAL])",
        )

    def test_extended_totals_row_also_keeps_structured_subtotal_reference(self):
        wb = _build_template_workbook()
        extended_classes = BASE_PAYMENT_CLASSES + ["11_new_method"]
        data = _make_data(extended_classes)
        report.write_by_payment_product_sheet(wb[report.BY_PAYMENT_PRODUCT_SHEET], data)

        ws = wb[report.BY_PAYMENT_PRODUCT_SHEET]
        self.assertEqual(
            ws.cell(row=13, column=12).value,
            f"=SUBTOTAL(109,{report.BY_PAYMENT_PRODUCT_TABLE}[TOTAL])",
        )


class ExistingStructureRegressionTests(unittest.TestCase):
    """Test 5: everything around the fix still behaves as before."""

    def test_table_ref_and_autofilter_and_formats_unchanged(self):
        wb = _build_template_workbook()
        data = _make_data(BASE_PAYMENT_CLASSES)
        report.write_by_payment_sheet(wb[report.BY_PAYMENT_SHEET], data)
        report.write_by_payment_product_sheet(wb[report.BY_PAYMENT_PRODUCT_SHEET], data)

        ws1 = wb[report.BY_PAYMENT_SHEET]
        ws2 = wb[report.BY_PAYMENT_PRODUCT_SHEET]
        t1 = ws1.tables[report.BY_PAYMENT_TABLE]
        t2 = ws2.tables[report.BY_PAYMENT_PRODUCT_TABLE]
        self.assertEqual(t1.ref, "A1:B12")
        self.assertEqual(t1.autoFilter.ref, "A1:B11")
        self.assertEqual(t2.ref, "A1:L12")
        self.assertEqual(t2.autoFilter.ref, "A1:L11")
        for row in range(2, 13):
            self.assertEqual(ws1.cell(row=row, column=2).number_format, "#,##0")
            for col in range(2, 13):
                self.assertEqual(ws2.cell(row=row, column=col).number_format, "#,##0")

    def test_payment_class_display_order_is_numeric_ascending(self):
        wb = _build_template_workbook()
        data = _make_data(BASE_PAYMENT_CLASSES)
        report.write_by_payment_sheet(wb[report.BY_PAYMENT_SHEET], data)

        ws1 = wb[report.BY_PAYMENT_SHEET]
        labels = [ws1.cell(row=row, column=1).value for row in range(2, 12)]
        self.assertEqual(
            labels,
            [
                "credit(JCB/AMEX)",
                "credit(SBPS)",
                "credit(DINERS)",
                "credit(その他)",
                "docomo",
                "au",
                "softbank",
                "paypay",
                "rakuten",
                "suica",
            ],
        )

    def test_product_column_order(self):
        wb = _build_template_workbook()
        ws2 = wb[report.BY_PAYMENT_PRODUCT_SHEET]
        headers = [ws2.cell(row=1, column=col).value for col in range(2, 12)]
        self.assertEqual(
            headers,
            ["100pt", "310pt", "410pt", "520pt", "630pt", "1050pt", "2150pt", "3250pt", "5450pt", "10800pt"],
        )

    def test_table_extends_for_11th_and_12th_payment_class(self):
        wb = _build_template_workbook()
        extended_classes = BASE_PAYMENT_CLASSES + ["11_new_method", "12_second_new_method"]
        data = _make_data(extended_classes)
        report.write_by_payment_sheet(wb[report.BY_PAYMENT_SHEET], data)
        report.write_by_payment_product_sheet(wb[report.BY_PAYMENT_PRODUCT_SHEET], data)

        t1 = wb[report.BY_PAYMENT_SHEET].tables[report.BY_PAYMENT_TABLE]
        t2 = wb[report.BY_PAYMENT_PRODUCT_SHEET].tables[report.BY_PAYMENT_PRODUCT_TABLE]
        self.assertEqual(t1.ref, "A1:B14")
        self.assertEqual(t2.ref, "A1:L14")

    def test_unexpected_product_price_fails_closed_even_for_zero_fill_row(self):
        with self.assertRaises(report.PlusPointSalesReportError) as ctx:
            report.build_report_data(
                [
                    {"price": 100, "payment_class": "1_credit(JCB/AMEX)", "sum_price": 100},
                    {"price": 7777, "payment_class": "1_credit(JCB/AMEX)", "sum_price": 0},
                ]
            )
        self.assertEqual(ctx.exception.code, "unexpected_product_price")

    def test_null_price_fails_closed(self):
        with self.assertRaises(report.PlusPointSalesReportError) as ctx:
            report.build_report_data([{"price": None, "payment_class": "1_credit(JCB/AMEX)", "sum_price": 500}])
        self.assertEqual(ctx.exception.code, "unexpected_product_price")

    def test_malformed_payment_class_fails_closed(self):
        with self.assertRaises(report.PlusPointSalesReportError) as ctx:
            report.build_report_data([{"price": 100, "payment_class": "credit_no_number", "sum_price": 100}])
        self.assertEqual(ctx.exception.code, "unparseable_payment_class")

    def test_unexpected_null_payment_class_with_nonzero_sales_fails_closed(self):
        with self.assertRaises(report.PlusPointSalesReportError) as ctx:
            report.build_report_data([{"price": 100, "payment_class": None, "sum_price": 500}])
        self.assertEqual(ctx.exception.code, "unexpected_payment_class")


if __name__ == "__main__":
    unittest.main()
