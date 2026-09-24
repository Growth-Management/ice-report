"""Builds synthetic XLSX templates for tests -- via openpyxl for the parts
it understands, then augmented with hand-written OOXML parts openpyxl has no
model for (xl/connections.xml, xl/queryTables/*, customXml/*, queryTable
relationships) so tests can verify xlsx_package_writer actually preserves
them. Test-only: production code never builds a package this way."""

from __future__ import annotations

import zipfile
from pathlib import Path

from openpyxl import Workbook
from openpyxl.worksheet.table import Table, TableColumn

CONNECTIONS_XML = b"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<connections xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">
  <connection id="2" name="\xe8\xa9\xb1\xe3\x83\x87\xe3\x83\xbc\xe3\x82\xbf_\xe4\xbd\x9c\xe5\x93\x81\xe5\x88\xa5" type="5" refreshedVersion="8" background="1"/>
  <connection id="3" name="\xe8\xa9\xb1\xe3\x83\x87\xe3\x83\xbc\xe3\x82\xbf_\xe4\xbd\x9c\xe5\x93\x81\xe5\x88\xa5_2" type="5" refreshedVersion="8" background="1"/>
</connections>
"""

QUERY_TABLE_XML_TEMPLATE = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
    '<queryTable xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
    'name="{name}" connectionId="{connection_id}" autoFormatId="16" applyNumberFormats="0" '
    'applyBorderFormats="0" applyFontFormats="0" applyPatternFormats="0" applyAlignmentFormats="0" '
    'applyWidthHeightFormats="1"/>\n'
)

TABLE_RELS_TEMPLATE = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
    '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
    '<Relationship Id="rId1" '
    'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/queryTable" '
    'Target="../queryTables/{query_table_file}"/>'
    "</Relationships>\n"
)

CUSTOM_XML_ITEM = b"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<root xmlns="urn:test-custom-xml"><note>template-origin custom data</note></root>
"""

CUSTOM_XML_ITEM_PROPS = b"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<ds:datastoreItem ds:itemID="{7B9A2F1A-0000-0000-0000-000000000001}" xmlns:ds="http://schemas.openxmlformats.org/officeDocument/2006/customXml"/>
"""

CUSTOM_XML_ITEM_RELS = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
    '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
    '<Relationship Id="rId1" '
    'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/customXmlProps" '
    'Target="itemProps1.xml"/>'
    "</Relationships>\n"
).encode("utf-8")


def _add_detail_sheet(wb: Workbook, name: str, headers: tuple[str, ...], *, freeze: str | None = None):
    ws = wb.create_sheet(name)
    for col, header in enumerate(headers, start=1):
        ws.cell(row=3, column=col, value=header)
    for col in range(1, len(headers) + 1):
        ws.cell(row=4, column=col).number_format = "#,##0"
    last_col_letter = ws.cell(row=3, column=len(headers)).column_letter
    table = Table(displayName=f"table_{name}".replace("　", "_"), ref=f"A3:{last_col_letter}4")
    table.tableColumns = [TableColumn(id=i, name=h) for i, h in enumerate(headers, start=1)]
    ws.add_table(table)
    if freeze:
        ws.freeze_panes = freeze
    return ws


def build_app2_like_template(output_path: str | Path) -> Path:
    """A template shaped like the real APP_2 official template: サマリ +
    iOS/Android/全体 (the sheets jumpplus_ad_revenue_report writes) plus
    作品別/作品別_2 (Power-Query-backed, never written by this module) --
    with connections.xml/queryTables/customXml/table rels injected
    afterward, exactly the parts an openpyxl round trip would silently
    drop."""
    wb = Workbook()
    wb.remove(wb.active)

    summary = wb.create_sheet("サマリ")  # サマリ
    summary["A1"] = "広告売上"
    summary["B3"] = "OS区分"
    summary["C3"] = "広告売上"
    summary["B4"] = "総計"
    summary["C4"] = None

    headers_10 = (
        "コンテンツID_Raise",
        "コンテンツID",
        "コンテンツ名",
        "JDCN",
        "広告表示数",
        "作品名",
        "コミックスJDCN",
        "コミックス巻数",
        "タイトルID",
        "デジタルタイトル名",
    )
    _add_detail_sheet(wb, "iOS", headers_10, freeze="D4")
    _add_detail_sheet(wb, "Android", headers_10, freeze="D4")
    zentai_headers = headers_10[:5] + ("広告売上", "広告売上_原資50") + headers_10[5:]
    _add_detail_sheet(wb, "全体", zentai_headers, freeze="I4")  # 全体

    sakuhin = _add_detail_sheet(wb, "作品別", ("A", "B"))  # 作品別
    sakuhin2 = _add_detail_sheet(wb, "作品別_2", ("A", "B", "C"))  # 作品別_2

    out_path = Path(output_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)

    _inject_power_query_parts(out_path, sakuhin_display=f"table_{sakuhin.title}", sakuhin2_display=f"table_{sakuhin2.title}")
    return out_path


def _table_file_for_display_name(entries: dict[str, bytes], display_name: str) -> str:
    for name, data in entries.items():
        if name.startswith("xl/tables/table") and display_name.encode("utf-8") in data:
            return name.rsplit("/", 1)[1]
    raise AssertionError(f"no table part found for displayName {display_name!r}")


def _inject_power_query_parts(path: Path, *, sakuhin_display: str, sakuhin2_display: str) -> None:
    with zipfile.ZipFile(path) as zf:
        entries = {name: zf.read(name) for name in zf.namelist()}

    sakuhin_table_file = _table_file_for_display_name(entries, sakuhin_display)
    sakuhin2_table_file = _table_file_for_display_name(entries, sakuhin2_display)

    # Mark both tables as queryTable-backed, like the real template.
    for table_file, conn_id in ((sakuhin_table_file, "2"), (sakuhin2_table_file, "3")):
        xml = entries[f"xl/tables/{table_file}"].decode("utf-8")
        xml = xml.replace("<table ", f'<table tableType="queryTable" ', 1)
        entries[f"xl/tables/{table_file}"] = xml.encode("utf-8")

    entries["xl/connections.xml"] = CONNECTIONS_XML
    entries["xl/queryTables/queryTable1.xml"] = QUERY_TABLE_XML_TEMPLATE.format(
        name="話データ_作品別", connection_id=2
    ).encode("utf-8")
    entries["xl/queryTables/queryTable2.xml"] = QUERY_TABLE_XML_TEMPLATE.format(
        name="話データ_作品別_2", connection_id=3
    ).encode("utf-8")
    entries[f"xl/tables/_rels/{sakuhin_table_file}.rels"] = TABLE_RELS_TEMPLATE.format(
        query_table_file="queryTable1.xml"
    ).encode("utf-8")
    entries[f"xl/tables/_rels/{sakuhin2_table_file}.rels"] = TABLE_RELS_TEMPLATE.format(
        query_table_file="queryTable2.xml"
    ).encode("utf-8")
    entries["customXml/item1.xml"] = CUSTOM_XML_ITEM
    entries["customXml/itemProps1.xml"] = CUSTOM_XML_ITEM_PROPS
    entries["customXml/_rels/item1.xml.rels"] = CUSTOM_XML_ITEM_RELS

    content_types = entries["[Content_Types].xml"].decode("utf-8")
    extra_overrides = (
        '<Override PartName="/xl/connections.xml" '
        'ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.connections+xml"/>'
        '<Override PartName="/xl/queryTables/queryTable1.xml" '
        'ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.queryTable+xml"/>'
        '<Override PartName="/xl/queryTables/queryTable2.xml" '
        'ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.queryTable+xml"/>'
        '<Override PartName="/customXml/item1.xml" ContentType="application/xml"/>'
        '<Override PartName="/customXml/itemProps1.xml" '
        'ContentType="application/vnd.openxmlformats-officedocument.customXmlProperties+xml"/>'
    )
    content_types = content_types.replace("</Types>", extra_overrides + "</Types>")
    entries["[Content_Types].xml"] = content_types.encode("utf-8")

    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zf:
        for name, data in entries.items():
            zf.writestr(name, data)


POWER_QUERY_PARTS = (
    "xl/connections.xml",
    "xl/queryTables/queryTable1.xml",
    "xl/queryTables/queryTable2.xml",
    "customXml/item1.xml",
    "customXml/itemProps1.xml",
    "customXml/_rels/item1.xml.rels",
)
